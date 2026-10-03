"""
75_sizer_beam_reconciliation.py - tally the empty-beam ledger with each sizer, beam number by beam number.
Ledger = ERP view WV_EMPTY_BEAM_STOCK_V from 15-Feb-2023, with the sign of empty beams returned (IGP_EMPTY_BEAMS) corrected
(the view counts them +1; they are -1). For every beam number the ledger says a sizer still holds, where is that number now
in our own records (latest movement, yarn_fabric_data.json)?
  - on a loom / sized waiting / empty in mill / at another sizer  -> the beam is NOT at that sizer: an entry mistake
  - the same number shows "returned more than sent" at another sizer -> the return was entered under the wrong party
  - nowhere in our records                                         -> ASK THE SIZER (or check the mill floor)
Hiba Fabrics brings its own beams - skipped.  ERP read only.
python 75_sizer_beam_reconciliation.py --upload <google sheet id>
"""
import configparser
import importlib.util
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
_d = importlib.util.spec_from_file_location("dash", HERE / "40_business_dashboard.py"); dash = importlib.util.module_from_spec(_d); _d.loader.exec_module(dash)
OUT = HERE / "financial_statements" / "sizer_beam_reconciliation.xlsx"
NAVY = PatternFill("solid", fgColor="1F3B57"); W = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
T = Font(bold=True, size=14, color="1F3B57"); NOTE = Font(italic=True, color="4B4F5A")
FILL = {"ASK SIZER": "F8C9C9", "Really at sizer": "D9F2E3", "Entry mistake": "FFF1CC", "Wrong party": "FFE3B3"}


def main():
    cur = dash.connect(dash.load_config()).cursor()
    cfg = configparser.ConfigParser(); cfg.read(HERE / "report_config.ini")
    data = json.loads((Path(cfg["output"]["folder"]) / "yarn_fabric_data.json").read_text(encoding="utf-8"))
    where = defaultdict(list)
    for b in data["beams"]:
        where[str(b[1])[:5]].append((b[2], str(b[3]).split(" · ")[0], str(b[4])[:10]))
    net = defaultdict(int); last = {}; hist = defaultdict(list)
    for an, td, tt, code, tno, gp in cur.execute(
            "SELECT an, TRUNC(td), tt, beam_code, tno, gp_no FROM dostex.wv_empty_beam_stock_v ORDER BY td, tno"):
        if "HIBA" in (an or ""):
            continue
        net[(an, code)] += 1 if tt == "OGP_EMPTY_BEAMS" else -1
        hist[(an, code)].append(f"{td:%d-%b-%y} {'sent' if tt == 'OGP_EMPTY_BEAMS' else 'back'} #{tno}")
        if tt == "OGP_EMPTY_BEAMS":
            last[(an, code)] = (td, tno, gp)
    over = defaultdict(list)                       # beam -> sizers where it came back more often than sent
    for (an, code), v in net.items():
        if v < 0:
            over[code].append(an)
    rows, summary = [], defaultdict(lambda: defaultdict(int))
    for (an, code), v in sorted(net.items()):
        if v <= 0:
            continue
        w = where.get(code, [])
        at_this = [x for x in w if x[0].startswith("At sizer") and x[1] == an]
        others = [x for x in w if not (x[0].startswith("At sizer") and x[1] == an)]
        wrong_p = [p for p in over.get(code, []) if p != an]
        if at_this:
            status, action = "Really at sizer", "Sent recently - should come back filled"
        elif others:
            status, action = "Entry mistake", "Not at this sizer: our records show it " + "; ".join(f"{x[0]} {x[1]} ({x[2]})" for x in others)
        elif wrong_p:
            status, action = "Wrong party", f"Came back but was entered under {', '.join(wrong_p)} - correct the party on the gate pass"
        else:
            status, action = "ASK SIZER", "No trace in our records since - ask the sizer to confirm / count beams in the mill"
        rows.append((an, code, v, last[(an, code)][0], last[(an, code)][1], last[(an, code)][2], status, action, " | ".join(hist[(an, code)][-6:])))
        summary[an][status] += v
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    ws["A1"] = "EMPTY BEAMS AT SIZERS - ledger tally, beam by beam"; ws["A1"].font = T
    ws["A2"] = ("Ledger = ERP empty-beam report from 15-Feb-2023 with its sign error corrected. Each beam the ledger says a sizer holds "
                "is checked against where that beam number is now in our records. Read from the ERP, nothing changed."); ws["A2"].font = NOTE
    heads = ["Sizer", "Ledger says at sizer", "Really at sizer", "ASK SIZER", "Wrong party", "Entry mistake"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=4, column=j, value=h); c.font, c.fill = W, NAVY
    r = 4
    for an in sorted(summary):
        r += 1; s = summary[an]
        ws.append([an, sum(s.values()), s["Really at sizer"], s["ASK SIZER"], s["Wrong party"], s["Entry mistake"]])
    ws.column_dimensions["A"].width = 36
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws.cell(row=r + 2, column=1, value="Next: send each sizer its 'ASK SIZER' beam numbers to confirm; staff correct 'Wrong party' and 'Entry mistake' "
                                       "gate passes in the ERP (see the Beams tab).").font = NOTE
    d = wb.create_sheet("Beams")
    heads = ["Sizer", "Beam #", "Ledger count", "Last sent", "OGP entry #", "Gate pass #", "Status", "What to do", "Ledger history (latest 6)"]
    for j, h in enumerate(heads, 1):
        c = d.cell(row=1, column=j, value=h); c.font, c.fill = W, NAVY
    for i, row in enumerate(sorted(rows, key=lambda x: (x[0], ["ASK SIZER", "Really at sizer", "Wrong party", "Entry mistake"].index(x[6]), x[1])), 2):
        for j, v in enumerate(row, 1):
            c = d.cell(row=i, column=j, value=v)
            c.fill = PatternFill("solid", fgColor=FILL[row[6]])
            c.alignment = Alignment(wrap_text=j in (8, 9), vertical="top")
        d.cell(row=i, column=4).number_format = "dd-mmm-yy"
    for col, w_ in zip("ABCDEFGHI", [32, 9, 8, 11, 9, 9, 15, 60, 70]):
        d.column_dimensions[col].width = w_
    d.freeze_panes = "C2"
    o = wb.create_sheet("Returned more than sent")
    for j, h in enumerate(["Sizer", "Beam #", "Extra returns", "Ledger history"], 1):
        c = o.cell(row=1, column=j, value=h); c.font, c.fill = W, NAVY
    k = 1
    for (an, code), v in sorted(net.items()):
        if v < 0:
            k += 1
            for j, val in enumerate([an, code, -v, " | ".join(hist[(an, code)][-6:])], 1):
                o.cell(row=k, column=j, value=val)
    for col, w_ in zip("ABCD", [32, 9, 12, 90]):
        o.column_dimensions[col].width = w_
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    for an in sorted(summary):
        print(an, dict(summary[an]))
    print("ASK SIZER:", [(r_[0], r_[1]) for r_ in rows if r_[6] == "ASK SIZER"])


def upload(fid):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    creds = service_account.Credentials.from_service_account_file(str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
    build("drive", "v3", credentials=creds).files().update(fileId=fid, media_body=MediaFileUpload(
        str(OUT), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
    print("uploaded", fid)


if __name__ == "__main__":
    main()
    if "--upload" in sys.argv:
        upload(sys.argv[sys.argv.index("--upload") + 1])
