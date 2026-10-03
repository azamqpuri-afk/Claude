"""
69_accountant_52ctn_marked.py - a copy of the accountant's 52 CTN closing workbook with:
  * a sheet "CLAUDE CHECK" in his layout - his figure and Claude's (ERP) figure side by side, difference, status, remark;
  * a sheet "ERRORS LIST" - every error / difference found, with the cell to look at and its effect in bags;
  * his own cells marked (red = error, orange = check / label, blue = method difference) with a note on each.
External-link formulas ([1]YARN!H33 ...) are replaced by their values so the copy opens anywhere. ERP read only.
python 69_accountant_52ctn_marked.py "<52 CTN.xlsx>" [--upload <google sheet id>]
"""
import importlib.util
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
_m = importlib.util.spec_from_file_location("a68", HERE / "68_audit_52ctn_workbook.py"); a68 = importlib.util.module_from_spec(_m); _m.loader.exec_module(a68)
OUT = HERE / "financial_statements" / "52 CTN closing - accountant vs Claude.xlsx"
RED = PatternFill("solid", fgColor="F8C9C9"); ORANGE = PatternFill("solid", fgColor="FFE3B3"); BLUE = PatternFill("solid", fgColor="D6E4F5")
GREEN = PatternFill("solid", fgColor="D9F2E3"); GREY = PatternFill("solid", fgColor="E8EDF2"); YEL = PatternFill("solid", fgColor="FFF4C2")
B = Font(bold=True); H = Font(bold=True, size=14); THIN = Side(style="thin", color="9AA5B1"); BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
N2 = '#,##0.00;[Red]-#,##0.00'; N0 = '#,##0;[Red]-#,##0'


def mark(ws, cell, fill, text):
    c = ws[cell]; c.fill = fill
    c.comment = Comment("CLAUDE CHECK: " + text, "Claude"); c.comment.width = 320; c.comment.height = 140


def main(path, upload_id=None):
    wb = load_workbook(path); val = load_workbook(path, data_only=True)
    for ws in wb.worksheets:                     # external links -> values
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "[" in c.value:
                    c.value = val[ws.title][c.coordinate].value
    v = val
    # ---- Claude's closing (55_) and beams by contract
    g = a68.gsheet(a68.CLOSING_55)
    cl = g["Closing"]; ol = g["On loom yarn"]
    d = lambda r, c="D": float(cl[f"{c}{r}"].value or 0)  # noqa: E731
    loom_by, wait_by, section = defaultdict(float), defaultdict(float), "loom"
    for row in ol.iter_rows(min_row=3, values_only=True):
        if row[0] and str(row[0]).startswith("Sized beams waiting"):
            section = "wait"
        if row[0] and str(row[0]).startswith("Yarn at the sizer"):
            break
        if isinstance(row[0], (int, float)) and isinstance(row[8], (int, float)):
            (loom_by if section == "loom" else wait_by)[int(row[0])] += float(row[8])
    mine = dict(recv=d(29), disp=-d(30), fresh=-d(31), loom=-d(32, "B"), unent=-d(32, "C"), wait=-d(33), siz=-d(34), hand=-d(36), bal=d(37),
                fresh_m=d(22, "B"), stock_m=d(20, "B"), rej_m=-d(21, "B"), bal60=float(cl["J31"].value or 0))
    s68, s70, S = v["68"], v["70"], v["SUMMARY"]
    recv = sum(-float(v[s][c].value) for s, c in (("38", "N74"), ("50", "N65"), ("60", "N56"), ("68", "N18"), ("70", "N14"))) / 100
    cons = sum(float(v[s][c].value) for s, c in (("38", "N75"), ("50", "N66"), ("60", "N57"), ("68", "N19"), ("70", "N15"))) / 100
    acc = dict(recv=recv, disp=cons, fresh=float(s68["T27"].value) + float(s68["U27"].value), loom=0.0, wait=float(s68["U29"].value or 0) + float(s70["U25"].value or 0),
               unent=0.0, siz=float(S["C8"].value), hand=float(s70["T26"].value or 0), bal=-float(S["C9"].value), bal60=-float(v["38"]["N82"].value))
    stock_m_acc = float(v["FABRIC STOCK"]["M37"].value) + float(v["REJECTION"]["M37"].value)
    ends_eff = (a68.WP - a68.WP90) * d(18, "B") / 100

    # ================= CLAUDE CHECK - his layout, two columns side by side
    ws = wb.create_sheet("CLAUDE CHECK", 0)
    ws["B1"] = "DOST TEXTILE"; ws["B1"].font = H
    ws["B2"] = "52/54 CTN CONTRACTS 38, 50, 60, 68, 70 - CLOSING CHECKED (ACCOUNTANT vs CLAUDE / ERP)"; ws["B2"].font = B
    ws["B3"] = f"AS ON 29/30 SEPT, 2026 · 1 BAG = 100 LBS · - = RETURNABLE (PAYABLE) TO QADRI, same sign as SUMMARY · checked {date.today():%d-%b-%Y}"
    heads = ["", "ITEM", "ACCOUNTANT (BAGS)", "CLAUDE - ERP (BAGS)", "DIFFERENCE (BAGS)", "STATUS", "REMARKS"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=5, column=i, value=h); c.font = B; c.fill = GREY; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
    L = [
        ("YARN RECEIVED 54/1 CTN (ALL CONTRACTS)", -acc["recv"], -mine["recv"], "OK", "Same on every contract: 1,820 + 1,619 + 1,983 + 253 + 650 bags = ERP IGPs less returns."),
        ("TOTAL YARN CONSUMED (FABRIC DISPATCHED)", acc["disp"], mine["disp"], "METHOD",
         f"Meters are the same. Ends: accountant 90, owner's closing 90.43 (actual sets) = {ends_eff:,.1f} bags. 60/1-weft fabric (15,819 m) kept separate in both."),
        ("FABRIC STOCK (FRESH) - ALL CONTRACTS, put on sheet 68", acc["fresh"], mine["fresh"], "ERROR",
         f"All contracts' fabric stock is on sheet 68 (fine for the total). Accountant: {stock_m_acc:,.0f} m total with 38,550 m rejection. ERP stock {mine['stock_m']:,.0f} m, owner's rejection {mine['rej_m']:,.0f} m -> fresh {mine['fresh_m']:,.0f} m."),
        # the accountant put all beam stock (on looms + waiting) in the SIZED BEAM rows of the last two sheets and left ON LOOM
        # blank on every sheet - so the two are compared together, per contract
        ("BEAMS ON LOOMS + SIZED BEAMS - CONTRACT 68 (sheet 68 U29)", float(s68["U29"].value or 0), loom_by[68] + wait_by[68], "ERROR",
         f"Accountant 18 in SIZED BEAM STOCK, ON LOOM blank. ERP: 9 beams on looms = {loom_by[68]:,.1f} bags, no waiting beams."),
        ("BEAMS ON LOOMS + SIZED BEAMS - CONTRACT 70 (sheet 70 U25)", float(s70["U25"].value or 0), loom_by[70] + wait_by[70], "ERROR",
         f"Accountant 185 in SIZED BEAM STOCK, ON LOOM blank. ERP: 10 beams on looms {loom_by[70]:,.1f} + 7 waiting beams {wait_by[70]:,.1f} = {loom_by[70] + wait_by[70]:,.1f} bags."),
        ("WEFT IN FABRIC WOVEN, NOT YET ENTERED", acc["unent"], mine["unent"], "ESTIMATE", "Claude's estimate (days since last production entry x loom's meters a day). Not in the accountant's closing."),
        ("STOCK AT SIZING (HAJI)", acc["siz"], mine["siz"], "OK", "Same 67 bags. Claude counts them as the fabric they will make (+1 bag)."),
        ("WEFT IN HAND", acc["hand"], mine["hand"], "ERROR", "Owner's count 68 bags on 28-Sep + IGP 595 of 29-Sep (130 bags) - issued to looms since = 194.4 (ERP book 190.9)."),
    ]
    r = 6
    for lab, a, m, st, rem in L:
        ws.cell(row=r, column=2, value=lab)
        ws.cell(row=r, column=3, value=round(a, 2)).number_format = N2
        ws.cell(row=r, column=4, value=round(m, 2)).number_format = N2
        ws.cell(row=r, column=5, value=f"=D{r}-C{r}").number_format = N2
        c = ws.cell(row=r, column=6, value=st); c.fill = {"OK": GREEN, "ERROR": RED, "METHOD": BLUE, "ESTIMATE": YEL}[st]; c.font = B
        ws.cell(row=r, column=7, value=rem)
        for col in range(2, 8):
            ws.cell(row=r, column=col).border = BOX; ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    # total: received (negative) + everything held / consumed = P/A (his sign)
    ws.cell(row=r, column=2, value="TOTAL P/A IN BAGS (- = PAYABLE TO QADRI)").font = B
    for col, L_ in ((3, "C"), (4, "D")):
        c = ws.cell(row=r, column=col, value=f"=SUM({L_}6:{L_}{r - 1})"); c.number_format = N2; c.font = B; c.fill = GREEN
    ws.cell(row=r, column=5, value=f"=D{r}-C{r}").number_format = N2
    ws.cell(row=r, column=7, value=f"Accountant's SUMMARY C9 = {-acc['bal']:,.2f} · Claude's closing = {-mine['bal']:,.2f}")
    for col in range(2, 8):
        ws.cell(row=r, column=col).border = BOX; ws.cell(row=r, column=col).font = B
    tot_r = r
    r += 2
    ws.cell(row=r, column=2, value="60/1 CTN WEFT (SEPARATE - TRANSFER FROM CONT#40)").font = B; r += 1
    ws.cell(row=r, column=2, value="YARN RETURNABLE IN BAGS (60/1)")
    ws.cell(row=r, column=3, value=round(-acc["bal60"], 3)); ws.cell(row=r, column=4, value=round(-mine["bal60"], 3))
    ws.cell(row=r, column=5, value=f"=D{r}-C{r}"); ws.cell(row=r, column=6, value="OK").fill = GREEN
    ws.cell(row=r, column=7, value="15,819 m at count 59 = 13.83 bags of the 14 transferred. Not carried to the accountant's SUMMARY (0.18 bag). ERP shows only 3.95 bags of 60/1 issued on contract 38.")
    r += 2
    ws.cell(row=r, column=2, value="WHERE THE 65 BAGS COME FROM").font = B; r += 1
    for lab, x in (("Accountant's errors (fabric stock, beams, weft in hand)", f"=E8+E9+E10+E13"),
                   ("Ends 90 vs 90.43 (method - owner's choice)", "=E7"),
                   ("Claude's estimate of unentered weft", "=E11"),
                   ("Sizing conversion", "=E12"),
                   ("TOTAL DIFFERENCE", f"=E{tot_r}")):
        ws.cell(row=r, column=2, value=lab); c = ws.cell(row=r, column=5, value=x); c.number_format = N2
        if lab.startswith("TOTAL"):
            ws.cell(row=r, column=2).font = B; c.font = B
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Colours in the accountant's sheets: RED = error · ORANGE = label / check · BLUE = method difference. Hover a coloured cell for the note.").font = Font(italic=True)
    for col, w in zip("ABCDEFG", [3, 46, 17, 17, 17, 12, 110]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C6"

    # ================= ERRORS LIST
    e = wb.create_sheet("ERRORS LIST", 1)
    e["B1"] = "DOST TEXTILE"; e["B1"].font = H
    e["B2"] = "52/54 CTN CLOSING - ERRORS AND DIFFERENCES FOUND (checked against the ERP)"; e["B2"].font = B
    for i, h in enumerate(["#", "SHEET / CELL", "WHAT THE SHEET SAYS", "CORRECT (ERP / OWNER)", "EFFECT (BAGS)", "TYPE", "FIX"], 1):
        c = e.cell(row=4, column=i, value=h); c.font = B; c.fill = GREY; c.border = BOX
    E = [
        ("68 U29 · 70 U25 (ON LOOM rows blank on all sheets)", "All beam stock put on the last two sheets: 18 (68) + 185 (70) = 203 bags",
         f"ERP beam by beam: contract 68 {loom_by[68]:,.1f} (on looms) · contract 70 {loom_by[70]:,.1f} on looms + {wait_by[70]:,.1f} waiting = {mine['loom'] + mine['wait']:,.1f}; "
         "contracts 38/50/60 have no beams left (0 is right)",
         mine["loom"] + mine["wait"] - acc["wait"], "ERROR", "Take the warp on each beam from the ERP (set length less shrinkage less woven) - see 'On loom yarn' in Claude's closing."),
        ("70 T26", "Weft in hand 185 bags", "194.4 bags (count 68 on 28-Sep + 130 bags IGP 595 on 29-Sep - issues)", mine["hand"] - acc["hand"], "ERROR", "Use the physical count, add receipts after it."),
        ("REJECTION M5 · FABRIC STOCK M5:M6", f"Total stock 148,413 m; rejection 38,550 m (balancing figure)",
         f"ERP stock {mine['stock_m']:,.0f} m; owner's rejection {mine['rej_m']:,.0f} m; fresh {mine['fresh_m']:,.0f} m", mine["fresh"] - acc["fresh"], "ERROR",
         "Take total stock from the ERP stock report; rejection = owner's count."),
        ("All sheets col B (ends)", "Ends 90 in every warp-bags formula", "Owner's closing uses 90.43 (actual sets: 12,480 ends)", ends_eff, "METHOD", "Owner to confirm which ends to use."),
        ("38 S69 (rejected)", "2,196 m rejected", "ERP returns 2,210.75 m: IGP 35 (15-Jan-25) 15 m 'rejection from Puri' not deducted; OGP 300 deducted 66 vs 65.75",
         None, "ERROR", "Deduct 14.75 m (Rs 562) from contract 38 billing."),
        ("50 S62 (rejected)", "934 m rejected", "ERP returns 973.75 m: IGP 1 (01-Dec-25) 40 m against OGP 477 not deducted; OGP 457 deducted 934 vs 933.75",
         None, "ERROR", "Deduct 39.75 m (Rs 1,684) from contract 50 billing."),
        ("50 (OGP list)", "OGP 468 not listed", "OGP 468 of 19-Oct-25: 2 m to HUZAIFA on contract 50", None, "ERROR", "Add the 2 m line (or confirm it is not Qadri's)."),
        ("38 N79", "60 CTN transfer 14 bags", "ERP shows only 3.95 bags of 60/1 issued on contract 38", None, "CHECK", "Book the transfer to contract 38 in the ERP or note where the rest was issued."),
        ("68 M18 · 70 M14", "'YARN RECEIVED 36/2 VP'", "54/1 CTN (the figure is cotton)", None, "LABEL", "Correct the label."),
        ("68 B3 · 70 B3", "'AS ON 15TH JULY, 2026'", "Data runs to 29-Sep-2026", None, "LABEL", "Update the date."),
        ("70 T26 / U25", "Weft in hand typed in WARP column; sized beams in WEFT column", "Sized beams = warp; yarn in hand = weft", None, "LABEL", "Swap the columns (total not affected)."),
        ("68 (whole sheet)", "Typed values, no formulas", "Other sheets use formulas", None, "CHECK", "Re-enter as formulas so it re-calculates."),
        ("68 T27/U27 · SUMMARY C5/C6", "All fabric stock booked to contract 68", "ERP stock is spread over 38, 60, 68, 70", None, "CHECK",
         "Per-contract P/A (68 +128.7, 70 -280.0) are not meaningful alone - only the total is."),
        ("38 N82", "60/1 balance -0.175 bag", "Not carried to SUMMARY", -0.175, "CHECK", "Add to SUMMARY (tiny)."),
    ]
    r = 5
    for i, (cell, says, right, eff, typ, fix) in enumerate(E, 1):
        vals = [i, cell, says, right, round(eff, 2) if isinstance(eff, float) else eff, typ, fix]
        for j, x in enumerate(vals, 1):
            c = e.cell(row=r, column=j, value=x); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
        e.cell(row=r, column=6).fill = {"ERROR": RED, "CHECK": ORANGE, "LABEL": ORANGE, "METHOD": BLUE}[typ]
        e.cell(row=r, column=5).number_format = N2
        r += 1
    for col, w in zip("ABCDEFG", [4, 28, 42, 60, 12, 9, 60]):
        e.column_dimensions[col].width = w
    e.freeze_panes = "A5"

    # ================= mark his own cells
    loomtxt = lambda c_: f"warp still on beams on looms, contract {c_}: {loom_by[c_]:,.1f} bags (ERP) - ON LOOM left blank"  # noqa: E731
    mark(wb["68"], "T28", ORANGE, loomtxt(68) + " here; the accountant's beam stock is in U29 (18), so compare the two together.")
    mark(wb["70"], "T24", ORANGE, loomtxt(70) + " here; the accountant's beam stock is in U25 (185), so compare the two together.")
    mark(wb["68"], "U29", RED, f"sized beams 18 typed in. ERP: contract 68 has no waiting beams; its beams on looms hold {loom_by[68]:,.1f} bags.")
    mark(wb["70"], "U25", RED, f"185 typed in. ERP: 7 waiting beams {wait_by[70]:,.1f} + on looms {loom_by[70]:,.1f} = {wait_by[70] + loom_by[70]:,.1f} bags. Also in the WEFT column - sized beams are warp.")
    mark(wb["70"], "T26", RED, "weft in hand should be 194.4 bags: count 68 on 28-Sep + 130 bags (IGP 595, 29-Sep) - issues since. ERP book 190.9. Also in the WARP column.")
    mark(wb["68"], "M18", ORANGE, "label says 36/2 VP - this is 54/1 cotton received on contract 68 (253 bags, matches ERP).")
    mark(wb["70"], "M14", ORANGE, "label says 36/2 VP - this is 54/1 cotton received on contract 70 (650 bags, matches ERP).")
    mark(wb["68"], "B3", ORANGE, "date not updated - OGPs and stock run to 29-Sep-2026."); mark(wb["70"], "B3", ORANGE, "date not updated - data runs to 29-Sep-2026.")
    mark(wb["REJECTION"], "M5", RED, f"148,413 typed as total stock; ERP stock is {mine['stock_m']:,.0f} m. Rejection (mill's) per owner is {mine['rej_m']:,.0f} m, not 38,550.")
    mark(wb["FABRIC STOCK"], "M37", RED, f"fresh stock should be ERP {mine['stock_m']:,.0f} - rejection {mine['rej_m']:,.0f} = {mine['fresh_m']:,.0f} m.")
    mark(wb["38"], "S69", RED, "ERP fabric returns 2,210.75 m: IGP 35 (15-Jan-25) 15 m 'rejection from Puri' not deducted; OGP 300 row deducted 66 vs 65.75 in ERP. Net 14.75 m still billed (Rs 562).")
    mark(wb["50"], "S62", RED, "ERP fabric returns 973.75 m: IGP 1 (01-Dec-25) 40 m against OGP 477 not deducted; OGP 457 row 934 vs 933.75. Net 39.75 m still billed (Rs 1,684).")
    mark(wb["50"], "Q62", RED, "OGP 468 (19-Oct-25, 2 m to HUZAIFA, contract 50) is in the ERP but not in this list.")
    mark(wb["38"], "N79", ORANGE, "14 bags 60/1 transfer - ERP shows only 3.95 bags of 60/1 issued on contract 38.")
    mark(wb["SUMMARY"], "C9", RED, f"Claude's closing (ERP): {-mine['bal']:,.2f} bags, {acc['bal'] - mine['bal']:,.2f} bags less payable to Qadri - see sheet CLAUDE CHECK.")
    for s_ in ("38", "50", "60", "68", "70"):
        mark(wb[s_], "B5", BLUE, "ends 90 used in warp bags; the owner's closing uses 90.43 ends (actual sets, 12,480 ends) - about 14 bags on all dispatched fabric.")
    OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
    print("saved", OUT, "acc", round(acc["bal"], 2), "mine", round(mine["bal"], 2), "loom_by", dict(loom_by), "wait_by", dict(wait_by))
    if upload_id:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload
        c = service_account.Credentials.from_service_account_file(str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
        build("drive", "v3", credentials=c).files().update(fileId=upload_id, media_body=MediaFileUpload(
            str(OUT), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
        print("uploaded", upload_id)


if __name__ == "__main__":
    a = sys.argv[1:]
    main(a[0], a[a.index("--upload") + 1] if "--upload" in a else None)
