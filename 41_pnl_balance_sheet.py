"""
41_pnl_balance_sheet.py
Project : Dost Textiles - Monthly P&L + yearly balance sheet, Jan-2014 to today
Purpose : Builds two .xlsx workbooks from the merged old + new ERP ledger
          (merged_ledger.py; old ERP DOSTT up to Sep-2023, new ERP DOSTEX from
          Oct-2023). Uploaded into the existing Google Sheets by
          upload_to_sheets() (service account imports the xlsx into the
          user-owned sheet, keeping the same link).
            dost_monthly_pnl.xlsx    - P&L per month + by fiscal year (Jul-Jun)
            dost_balance_sheet.xlsx  - balance sheet at every 30 June + the
                                       switch-over date + today; trial balance
                                       today; account mapping; old->new
                                       opening reconciliation
Conventions:
  - Sign: ledger AMT > 0 debit. Revenue, liabilities and equity are shown
    positive by flipping the sign.
  - Neither ERP ever closed a year: P&L accounts accumulate from Jan-2014, so
    equity = capital accounts + accumulated profit since Jan-2014.
  - Customers in credit (advances received) are shown as a liability and
    suppliers in debit (advances paid) as an asset, per party.
"""

import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import importlib.util  # noqa: E402

import closing_entries as ce  # noqa: E402
import merged_ledger as ml  # noqa: E402
from old_erp_account_map import HEAD_MAP, LEAF_OVERRIDES  # noqa: E402

_s2 = importlib.util.spec_from_file_location("stores", HERE / "45_stores_reconciliation_pdf.py")
_stores = importlib.util.module_from_spec(_s2)
_s2.loader.exec_module(_stores)
_spec = importlib.util.spec_from_file_location("dash", HERE / "40_business_dashboard.py")
dash = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(dash)

OUT_DIR = HERE / "financial_statements"
SHEETS = {"dost_monthly_pnl.xlsx": "1BHOkvTFqaDINZv_Kldpg-va00KIO0q1cjD3gz_h-U90",
          "dost_balance_sheet.xlsx": "1RIup5nTtrH-QF6FSz-ajtc4n5Lk3g4NwvqATZ6HKJI0"}

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3B57")
OLD_FILL = PatternFill("solid", fgColor="5B6B7A")
SEC_FONT = Font(bold=True, color="1F3B57")
SEC_FILL = PatternFill("solid", fgColor="E8EDF2")
TOT_FONT = Font(bold=True)
TOP = Border(top=Side(style="thin", color="7F8C99"))
DBL = Border(top=Side(style="thin", color="1F3B57"), bottom=Side(style="double", color="1F3B57"))
NUM = '#,##0;[Red](#,##0);"â€“"'
NOTE = Font(italic=True, size=9, color="555555")
WARN = Font(italic=True, size=9, color="9A3412")


def mon_label(m):
    return date(int(m[:4]), int(m[5:]), 1).strftime("%b-%y")


def fy_of(m):
    y, mo = int(m[:4]), int(m[5:])
    return f"FY{(y + 1 if mo >= 7 else y) % 100:02d}"


def header(ws, row, labels, old_flags=None):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=1 + i, value=h)
        c.font = HEAD
        c.fill = OLD_FILL if old_flags and old_flags[i] else HEAD_FILL
        c.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center", wrap_text=True)


def write_row(ws, r, code, name, values, kind=None):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=2, value=name)
    for i, v in enumerate(values):
        c = ws.cell(row=r, column=3 + i, value=None if v is None else round(v))
        c.number_format = NUM
    ncol = 2 + len(values)
    if kind == "section":
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).font = SEC_FONT
            ws.cell(row=r, column=c).fill = SEC_FILL
    elif kind in ("total", "grand"):
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).font = TOT_FONT
            ws.cell(row=r, column=c).border = DBL if kind == "grand" else TOP


def pct_row(ws, r, label, num, den):
    ws.cell(row=r, column=2, value=label)
    for i, (a, b) in enumerate(zip(num, den)):
        c = ws.cell(row=r, column=3 + i, value=(a / b) if b else None)
        c.number_format = "0.0%"


# ------------------------------------------------------------------ P&L
def build_pnl(L, names):
    months = L.months
    mi = {m: i for i, m in enumerate(months)}
    lines = defaultdict(lambda: [0.0] * len(months))   # key -> monthly amounts (natural sign)
    draws = [0.0] * len(months)                         # partners' drawings (treated as expense)
    for era, m, ac, an, h, vt, tbl, amt in L.records:
        if h == "100303" and m in mi:
            draws[mi[m]] += amt
        if h[:2] in ("35", "40", "45", "55") and m in mi:
            key = "350105002" if ac == "350105002" else h
            lines[key][mi[m]] += amt
    fys = []
    for m in months:
        if fy_of(m) not in fys:
            fys.append(fy_of(m))
    label = {"350105002": "CONVERSION YARN SALE INCOME"}

    wb = Workbook()
    views = [("Monthly P&L", months), ("P&L by year", fys)]
    for idx, (title, cols) in enumerate(views):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws["A1"] = "Dost Textiles - Profit & Loss " + ("(monthly)" if idx == 0 else "(fiscal year July-June)")
        ws["A1"].font = Font(bold=True, size=14, color="1F3B57")
        ws["A2"] = (f"Old ERP Jan-2014 to Sep-2023 (grey column headers) + new ERP Oct-2023 to {date.today():%d-%b-%Y}. "
                    "Rs. Current month/year is incomplete; month-end journals post late.")
        ws["A2"].font = NOTE
        if idx == 0:
            labels = ["Code", "Account"] + [mon_label(c) for c in cols] + ["Total"]
            flags = [False, False] + [ml.is_old(c) for c in cols] + [False]
            pick = lambda s: list(s) + [sum(s)]  # noqa: E731
        else:
            labels = ["Code", "Account"] + [c + (" (to date)" if c == fys[-1] else "") + (" *" if c == "FY24" else "")
                                            for c in cols] + ["Total"]
            flags = [False, False] + [c < "FY24" for c in cols] + [False]

            def pick(s):
                v = [sum(x for m, x in zip(months, s) if fy_of(m) == f) for f in fys]
                return v + [sum(v)]
        header(ws, 4, labels, flags)
        r = 5
        width = len(cols) + 1

        def section(t):
            nonlocal r
            write_row(ws, r, None, t, [None] * width, "section"); r += 1

        def block(prefix, sign, total_label):
            nonlocal r
            tot = [0.0] * width
            for key in sorted(k for k in lines if k.startswith(prefix)):
                v = [sign * x for x in pick(lines[key])]
                if any(abs(x) >= 1 for x in v):
                    write_row(ws, r, key, label.get(key, names.get(key, key)), v); r += 1
                    tot = [a + b for a, b in zip(tot, v)]
            write_row(ws, r, None, total_label, tot, "total"); r += 1
            return tot

        section("REVENUE")
        rev = block("35", -1, "Total revenue"); r += 1
        section("COST OF PRODUCTION (manufacturing expenses)")
        cop = block("40", 1, "Total cost of production"); r += 1
        gp = [a - b for a, b in zip(rev, cop)]
        write_row(ws, r, None, "GROSS PROFIT", gp, "grand"); r += 1
        pct_row(ws, r, "Gross margin %", gp, rev); r += 2
        section("ADMINISTRATIVE & SELLING EXPENSES")
        adm = block("45", 1, "Total admin & selling"); r += 1
        section("FINANCIAL EXPENSES")
        fin = block("55", 1, "Total financial expenses"); r += 1
        net = [a - b - c for a, b, c in zip(gp, adm, fin)]
        write_row(ws, r, None, "NET PROFIT / (LOSS)", net, "grand"); r += 1
        pct_row(ws, r, "Net margin %", net, rev); r += 2
        dr = pick(draws)
        write_row(ws, r, None, "Less: partners' drawings (treated as expense; not tax-deductible)", [-v for v in dr]); r += 1
        write_row(ws, r, None, "PROFIT AFTER DRAWINGS (before tax; split 50/50 after tax)", [a - b for a, b in zip(net, dr)], "grand"); r += 2
        ws.cell(row=r, column=2, value="No depreciation, raw material consumption or stock movement has been booked in "
                "either ERP, and the years were never closed; profit is before these and before tax.").font = WARN
        if idx == 1:
            ws.cell(row=r + 1, column=2, value="* FY24 = Jul-2023 to Jun-2024: Jul-Sep 2023 from the old ERP, "
                    "Oct-2023 onward from the new ERP. FY14 covers Jan-Jun 2014 only.").font = NOTE
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 40
        for i in range(len(labels) - 2):
            ws.column_dimensions[get_column_letter(3 + i)].width = 12 if idx == 0 else 14
        ws.freeze_panes = "C5"
    return wb


# ------------------------------------------------------------------ balance sheet
def build_bs(L, names, cur):
    today = date.today()
    dates = []
    for y in range(2014, today.year + 1):
        if date(y, 6, 30) < today:
            dates.append((date(y, 6, 30), f"30-Jun-{y % 100:02d}"))
    dates.append((date(2023, 9, 30), "30-Sep-23 (old ERP close)"))
    dates.sort()
    dates.append((today, f"{today:%d-%b-%y} (today)"))
    names = dict(names)
    names.update({"209001": "ADVANCES FROM CUSTOMERS", "309001": "ADVANCES TO SUPPLIERS"})

    bals = []
    for d, _ in dates:
        b = defaultdict(float)
        for (ac, h), v in L.leaf_balances_at(d.isoformat()).items():
            if h.startswith("3003") and v < 0:
                h = "209001"
            elif h.startswith("2001") and v > 0:
                h = "309001"
            b[h] += v
        bals.append(b)

    wb = Workbook()
    ws = wb.active
    ws.title = "Balance sheet"
    ws["A1"] = "Dost Textiles - Balance Sheet at 30 June each year (and at switch-over and today)"
    ws["A1"].font = Font(bold=True, size=14, color="1F3B57")
    ws["A2"] = ("Old ERP up to 30-Sep-2023 (grey column headers), new ERP after. Rs. No depreciation, stock valuation "
                "or tax has been booked in either ERP, so assets are at cost and profit is before these. Includes the "
                "owner-approved correction: YARN, THREAD STOCK (-27,631, FY25) moved to yarn sale income.")
    ws["A2"].font = NOTE
    header(ws, 4, ["Code", "Head"] + [lbl for _, lbl in dates], [False, False] + [d.isoformat() <= ml.OLD_LAST_DATE for d, _ in dates])
    r, n = 5, len(dates)

    def section(t):
        nonlocal r
        write_row(ws, r, None, t, [None] * n, "section"); r += 1

    def group(tops, sign, label):
        nonlocal r
        heads = sorted({h for b in bals for h in b if h[:2] in tops})
        tot = [0.0] * n
        for h in heads:
            v = [sign * b.get(h, 0.0) for b in bals]
            if any(abs(x) >= 1 for x in v):
                write_row(ws, r, h, names.get(h, h), v); r += 1
                tot = [a + b for a, b in zip(tot, v)]
        write_row(ws, r, None, label, tot, "total"); r += 1
        return tot

    section("ASSETS")
    nca = group(["25"], 1, "Total non-current assets")
    ca = group(["30"], 1, "Total current assets")
    ta = [a + b for a, b in zip(nca, ca)]
    write_row(ws, r, None, "TOTAL ASSETS", ta, "grand"); r += 2
    section("LIABILITIES")
    ltl = group(["15"], -1, "Total long-term liabilities")
    cl = group(["20"], -1, "Total current liabilities")
    tl = [a + b for a, b in zip(ltl, cl)]
    write_row(ws, r, None, "TOTAL LIABILITIES", tl, "grand"); r += 2
    section("EQUITY")
    cap = group(["10"], -1, "Capital, reserves & drawings")
    prof = [-sum(v for h, v in b.items() if h[:2] in ("35", "40", "45", "55")) for b in bals]
    write_row(ws, r, None, "Accumulated profit since Jan-2014 (years never closed)", prof); r += 1
    te = [a + b for a, b in zip(cap, prof)]
    write_row(ws, r, None, "TOTAL EQUITY", te, "grand"); r += 2
    tle = [a + b for a, b in zip(tl, te)]
    write_row(ws, r, None, "TOTAL LIABILITIES + EQUITY", tle, "grand"); r += 1
    diff = [a - b for a, b in zip(ta, tle)]
    write_row(ws, r, None, "Difference (assets - liabilities - equity)", diff)
    for c in range(1, n + 3):
        ws.cell(row=r, column=c).font = Font(bold=True, color="B91C1C")
    r += 1
    ws.cell(row=r, column=2, value="Any difference after Sep-2023 comes from store receipt (SRN) vouchers posted "
            "without their sales-tax line (see flag 1).").font = WARN
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 46
    for i in range(n):
        ws.column_dimensions[get_column_letter(3 + i)].width = 15
    ws.freeze_panes = "C5"

    old_flags = [d.isoformat() <= ml.OLD_LAST_DATE for d, _ in dates]
    closed_diff = ce.build_closing_sheets(
        wb, L, names, dates, (header, write_row, NUM, NOTE, WARN, TOT_FONT, get_column_letter, old_flags),
        ce.load_srn(cur, dash.rows), _stores.load_store_gap(cur, dash.rows, from_month="2016-10"))
    print("closed diff", [round(x) for x in closed_diff])

    # Trial balance today (new ERP leaves)
    tb = wb.create_sheet("Trial balance today")
    tb["A1"] = f"Trial balance as of {today:%d-%b-%Y} (new ERP leaf accounts, incl. balances brought forward)"
    tb["A1"].font = Font(bold=True, size=13, color="1F3B57")
    header(tb, 3, ["Code", "Account", "Head", "Debit", "Credit"])
    r = 4
    dr_t = cr_t = 0.0
    for (ac, h), v in sorted(L.leaf_balances_at(today.isoformat()).items()):
        if abs(v) < 1:
            continue
        tb.cell(row=r, column=1, value=ac)
        tb.cell(row=r, column=2, value=names.get(ac, ac))
        tb.cell(row=r, column=3, value=names.get(h, h))
        c = tb.cell(row=r, column=4 if v > 0 else 5, value=round(abs(v)))
        c.number_format = NUM
        dr_t += max(v, 0); cr_t += max(-v, 0)
        r += 1
    tb.cell(row=r, column=2, value="Total").font = TOT_FONT
    for col, v in ((4, dr_t), (5, cr_t)):
        c = tb.cell(row=r, column=col, value=round(v))
        c.number_format, c.font, c.border = NUM, TOT_FONT, DBL
    for col, wdt in zip("ABCDE", (12, 44, 34, 15, 15)):
        tb.column_dimensions[col].width = wdt
    tb.freeze_panes = "A4"

    # Account mapping old -> new
    mp = wb.create_sheet("Old to new account map")
    mp["A1"] = "How old-ERP heads were mapped onto the new chart of accounts"
    mp["A1"].font = Font(bold=True, size=13, color="1F3B57")
    header(mp, 3, ["Old head", "Old head name", "New head", "New head name", "Old balance 30-Sep-23"])
    old_names = dict(dash.rows(cur, "SELECT ac, an FROM dostt.ach"))
    old_close = defaultdict(float)
    for era, m, ac, an, h, vt, tbl, amt in L.records:
        if era == "old":
            old_close[ac[:5]] += amt
    r = 4
    for oh in sorted(HEAD_MAP):
        mp.cell(row=r, column=1, value=oh)
        mp.cell(row=r, column=2, value=(old_names.get(oh) or "").strip())
        mp.cell(row=r, column=3, value=HEAD_MAP[oh])
        mp.cell(row=r, column=4, value=names.get(HEAD_MAP[oh], ""))
        c = mp.cell(row=r, column=5, value=round(old_close.get(oh, 0.0)))
        c.number_format = NUM
        r += 1
    r += 1
    mp.cell(row=r, column=1, value="Individual accounts moved to a different head (the new chart splits these):").font = TOT_FONT
    r += 1
    for leaf, nh in sorted(LEAF_OVERRIDES.items()):
        mp.cell(row=r, column=1, value=leaf)
        mp.cell(row=r, column=2, value=(old_names.get(leaf) or "").strip())
        mp.cell(row=r, column=3, value=nh)
        mp.cell(row=r, column=4, value=names.get(nh, ""))
        r += 1
    for col, wdt in zip("ABCDE", (12, 40, 12, 40, 20)):
        mp.column_dimensions[col].width = wdt

    # Reconciliation old close vs new opening
    rc = wb.create_sheet("Old close vs new opening")
    rc["A1"] = "Old ERP closing balances (30-Sep-2023) vs new ERP opening balances, by new head"
    rc["A1"].font = Font(bold=True, size=13, color="1F3B57")
    rc["A2"] = ("Differences are reclassifications made when the new ERP was set up; they net to zero overall.")
    rc["A2"].font = NOTE
    header(rc, 4, ["Head", "Name", "Old close", "New opening", "Difference"])
    oc = defaultdict(float)
    for (ac, h), v in L.leaf_balances_at(ml.OLD_LAST_DATE).items():
        oc[h] += v
    no = defaultdict(float)
    for ac, an, h, v in L.new_opening:
        no[h] += v
    r = 5
    td = 0.0
    for h in sorted(set(oc) | set(no)):
        a, b = oc.get(h, 0.0), no.get(h, 0.0)
        if abs(a) < 1 and abs(b) < 1:
            continue
        rc.cell(row=r, column=1, value=h)
        rc.cell(row=r, column=2, value=names.get(h, h))
        for col, v in ((3, a), (4, b), (5, a - b)):
            c = rc.cell(row=r, column=col, value=round(v))
            c.number_format = NUM
        if abs(a - b) >= 1:
            for col in range(1, 6):
                rc.cell(row=r, column=col).font = Font(color="9A3412")
        td += a - b
        r += 1
    rc.cell(row=r, column=2, value="Total difference").font = TOT_FONT
    c = rc.cell(row=r, column=5, value=round(td))
    c.number_format, c.font = NUM, TOT_FONT
    for col, wdt in zip("ABCDE", (10, 40, 16, 16, 16)):
        rc.column_dimensions[col].width = wdt
    return wb, dates, diff


def upload_to_sheets():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    creds = service_account.Credentials.from_service_account_file(
        str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
    drive = build("drive", "v3", credentials=creds)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    for fname, fid in SHEETS.items():
        drive.files().update(fileId=fid, media_body=MediaFileUpload(str(OUT_DIR / fname), mimetype=xlsx)).execute()
        print("uploaded", fname, "->", fid)


def main():
    cfg = dash.load_config()
    con = dash.connect(cfg)
    cur = con.cursor()
    months = dash.month_list(ml.FIRST_MONTH, date.today().strftime("%Y-%m"))
    L = ml.Ledger(cur, dash.rows, months)
    names = L.names
    OUT_DIR.mkdir(exist_ok=True)
    build_pnl(L, names).save(OUT_DIR / "dost_monthly_pnl.xlsx")
    bs, dates, diff = build_bs(L, names, cur)
    bs.save(OUT_DIR / "dost_balance_sheet.xlsx")
    print("bs dates", [d.isoformat() for d, _ in dates])
    print("diff", [round(x) for x in diff])
    if "--upload" in sys.argv:
        upload_to_sheets()


if __name__ == "__main__":
    main()



