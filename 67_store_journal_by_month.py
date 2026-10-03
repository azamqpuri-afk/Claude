"""
67_store_journal_by_month.py - monthly store consumption journal ("BILL M/O ... (STORE)") vs the store's issue and
return slips, one tab per ERP (old DOSTT Oct-2016 to Sep-2023, new DOSTEX from Oct-2023), newest month first.
Difference = journal - (issues - returns); minus = the journal charged less than the store used. ERP read only.
Upload:  python 67_store_journal_by_month.py --upload <google sheet id>
"""
import importlib.util
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
_s = importlib.util.spec_from_file_location("st", HERE / "45_stores_reconciliation_pdf.py"); st = importlib.util.module_from_spec(_s); _s.loader.exec_module(st)
_d = importlib.util.spec_from_file_location("d", HERE / "40_business_dashboard.py"); dash = importlib.util.module_from_spec(_d); _d.loader.exec_module(dash)
OUT = HERE / "financial_statements" / "store_journal_vs_slips_by_month.xlsx"
NAVY = PatternFill("solid", fgColor="1F3B57"); RED = PatternFill("solid", fgColor="FBE3E3"); GREEN = PatternFill("solid", fgColor="D9F2E3")
W = Font(bold=True, color="FFFFFF"); B = Font(bold=True); T = Font(bold=True, size=14, color="1F3B57"); NOTE = Font(italic=True, color="4B4F5A")
N0 = '#,##0;[Red]-#,##0'
# months with their own suggested entry in the corrections PDF (A5)
LISTED = {"2023-10", "2025-01", "2025-07", "2025-11", "2025-12", "2026-06", "2026-07"}


def tab(ws, title, months, iss, ret, jv, why, note, last):
    ws["A1"] = title; ws["A1"].font = T
    ws["A2"] = note; ws["A2"].font = NOTE
    heads = ["Month", "JV #", "Journal narration", "Issue slips (Rs)", "Return slips (Rs)", "Store net (Rs)",
             "Journal - credit to Stores & Spares Inventory (Rs)", "Difference (Rs)", "Likely reason", "In the corrections PDF"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=4, column=i, value=h); c.font, c.fill = W, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    r = 5
    for m in sorted(months, reverse=True):
        vno, nar, j = jv.get(m, ("", "", 0.0))
        ws.cell(row=r, column=1, value=datetime.strptime(m, "%Y-%m")).number_format = "mmm-yyyy"
        ws.cell(row=r, column=2, value=str(vno or "not posted"))
        ws.cell(row=r, column=3, value=nar or "")
        ws.cell(row=r, column=4, value=round(iss[m], 2))
        ws.cell(row=r, column=5, value=round(ret[m], 2))
        ws.cell(row=r, column=6, value=f"=D{r}-E{r}")
        ws.cell(row=r, column=7, value=round(j, 2))
        # not compared: the open month (journal not posted yet) and old months before the store module (no slips)
        skip = m > last or m < "2016-10"
        ws.cell(row=r, column=8, value=None if skip else f"=G{r}-F{r}")
        ws.cell(row=r, column=9, value=("Journal not posted yet (month still open) - not compared" if m > last else
                                        "Before the store module started (no slips) - not compared" if m < "2016-10" else why.get(m, "")))
        ws.cell(row=r, column=10, value="" if skip else ("own entry" if m in LISTED else ("in net entry 13" if m >= "2023-10" else "in old-ERP net entry")))
        for col in range(4, 9):
            ws.cell(row=r, column=col).number_format = N0
        d = j - (iss[m] - ret[m])
        if abs(d) >= 50000 and not skip:
            for col in range(1, 11):
                ws.cell(row=r, column=col).fill = RED
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = B
    for col in "DEFGH":
        # only the compared months (Difference filled), so the total row adds up across
        c = ws[f"{col}{r}"]; c.value = f'=SUMIF($H$5:$H${r - 1},"<>",{col}5:{col}{r - 1})'; c.font = B; c.number_format = N0; c.fill = GREEN
    for col, w in zip("ABCDEFGHIJ", [10, 10, 30, 14, 13, 14, 20, 14, 70, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    return r


def return_tab(ws, cur):
    """The ERP has no Store Return report (owner, 1-Oct-2026): monthly totals + every return slip line (new ERP), valued qty x average rate."""
    lines = dash.rows(cur, """SELECT TO_CHAR(m.dd,'YYYY-MM'), m.dno, TRUNC(m.dd), d.ic, i.idesc, i.uom, d.qty, NVL(d.avg_rate, d.rate),
                                     ROUND(d.qty * NVL(d.avg_rate, d.rate), 2), m.feedd, m.remarks
                              FROM dostex.dmn m JOIN dostex.ddt d ON d.trans_no = m.trans_no AND d.uc = m.uc
                              JOIN dostex.i i ON i.ic = d.ic AND i.uc = d.uc
                              WHERE m.dt = 'RET' AND m.dd >= DATE '2023-10-01' ORDER BY m.dd DESC, m.dno DESC, d.sr_no""")
    ws["A1"] = "STORE RETURN SLIPS (new ERP) - the ERP has no return report, so here is every return slip, newest first"; ws["A1"].font = T
    ws["A2"] = ("Value = quantity x average rate, as the ERP values them today (average rate can move when a GRN is entered late). "
                "The store journal's credit to 300101001 should be: Issue slips - this month's Return total."); ws["A2"].font = NOTE
    tot = {}
    for ln in lines:
        tot[ln[0]] = tot.get(ln[0], 0.0) + float(ln[8] or 0)
    r = 4
    for j, hd in enumerate(["Month", "Return slips total (Rs)", "Slips", "Lines"], 1):
        c = ws.cell(row=r, column=j, value=hd); c.font, c.fill = W, NAVY
    for mo in sorted(tot, reverse=True):
        r += 1
        ws.cell(row=r, column=1, value=datetime.strptime(mo, "%Y-%m")).number_format = "mmm-yyyy"
        ws.cell(row=r, column=2, value=round(tot[mo], 2)).number_format = N0
        ws.cell(row=r, column=3, value=len({ln[1] for ln in lines if ln[0] == mo}))
        ws.cell(row=r, column=4, value=sum(1 for ln in lines if ln[0] == mo))
    r += 2
    heads = ["Month", "Return #", "Return date", "Item code", "Item", "Unit", "Qty", "Average rate", "Value (Rs)", "Entered on", "Remarks", "Check"]
    for j, hd in enumerate(heads, 1):
        c = ws.cell(row=r, column=j, value=hd); c.font, c.fill = W, NAVY
    for mo, dno, dd, ic, idesc, uom, qty, rate, val, fd, rem in lines:
        r += 1
        chk = ""
        if rate is None:
            chk = "no average rate - value 0 (item had no stock value when returned)"
        elif float(val or 0) >= 100000:
            chk = "very large value for a return - check the unit / quantity"
        row = [datetime.strptime(mo, "%Y-%m"), dno, dd, ic, idesc, uom, float(qty or 0), float(rate) if rate is not None else None, float(val or 0), fd, rem or "", chk]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            if chk:
                c.fill = RED
        ws.cell(row=r, column=1).number_format = "mmm-yyyy"; ws.cell(row=r, column=3).number_format = "dd-mmm-yy"
        ws.cell(row=r, column=10).number_format = "dd-mmm-yy hh:mm"; ws.cell(row=r, column=9).number_format = N0
    for col, w in zip("ABCDEFGHIJKL", [10, 10, 11, 12, 42, 7, 8, 12, 13, 15, 30, 50]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"


def main():
    cur = dash.connect(dash.load_config()).cursor()
    months, iss, ret, jv = st.monthly(cur, dash.rows)
    last = max(jv)
    diff, flagged = st.reasons(months, iss, ret, jv, last)
    why = {m: v[1] for m, v in flagged.items()}
    for m in months:
        if m not in why and abs(diff[m]) >= 1:
            net, j = iss[m] - ret[m], jv.get(m, (0, 0, 0.0))[2]
            if ret[m] and abs(j - iss[m]) < 1:
                why[m] = "Return slips not deducted (journal = issue slips)"
            elif ret[m] and abs(diff[m] - ret[m]) < 0.2 * ret[m]:
                why[m] = "Return slips probably not deducted (difference close to the returns)"
    wb = Workbook()
    new = [m for m in months if m >= "2023-10"]
    old = [m for m in months if m < "2023-10"]
    tab(wb.active, "NEW ERP (from Oct-2023) - store consumption journal vs store slips, newest first", new, iss, ret, jv, why,
        "Difference = journal - (issue slips - return slips), slips valued at average rate. Minus = journal charged less than the store used "
        "(consumption too low, inventory too high). Red rows differ by Rs 50,000 or more. Read from the ERP, nothing changed.", last)
    wb.active.title = "New ERP"
    tab(wb.create_sheet("Old ERP"), "OLD ERP (to Sep-2023) - store consumption journal vs store slips, newest first", old, iss, ret, jv, why,
        "Same check on the old ERP (DOSTT). The old ERP is closed, so its net difference is corrected once in the new ERP. "
        "Months before the store module started have a journal but no slips to compare.", last)
    h = wb.create_sheet("How to verify")
    for i, t in enumerate([
        "HOW TO VERIFY ANY MONTH IN THE ERP",
        "1. Store Issue report for the month (1st to last day): total VALUE AT AVERAGE RATE = 'Issue slips'.",
        "2. Return slips: the ERP has no Store Return report - use the 'Return slips' tab (monthly total + every slip line, valued at average rate like the ERP).",
        "3. Open the month's JV (JV # column, narration 'BILL M/O ... (STORE)'): the CREDIT to 300101001 STORES & SPARES-INVENTORY = 'Journal'. "
        "Or see the 300101001 ledger for that month.",
        "4. Step 3 should equal step 1 - step 2. The gap is the Difference.",
        "Use values at average rate - the ERP values stock at average rate; a report at purchase rate will differ slightly.",
        "Correct: Difference minus -> Dr 400301001 STORES AND SPARES CONSUMED / Cr 300101001 STORES & SPARES-INVENTORY; plus -> the reverse.",
    ], 1):
        h.cell(row=i, column=1, value=t).font = T if i == 1 else Font()
    h.column_dimensions["A"].width = 150
    return_tab(wb.create_sheet("Return slips", 1), cur)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    sm = lambda ms: sum(jv.get(m, (0, 0, 0.0))[2] - (iss[m] - ret[m]) for m in ms)  # noqa: E731
    print("new", len(new), round(sm(new)), "small-only", round(sm([m for m in new if m not in LISTED and m <= last])), "| old", len(old), round(sm(old)),
          "old Oct-16..Sep-23", round(sm([m for m in old if m >= "2016-10"])))


def upload(fid):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    creds = service_account.Credentials.from_service_account_file(str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
    build("drive", "v3", credentials=creds).files().update(fileId=fid, media_body=MediaFileUpload(
        str(OUT), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
    print("uploaded", fid)


if __name__ == "__main__":
    main()
    if len(sys.argv) > 2 and sys.argv[1] == "--upload":
        upload(sys.argv[2])
