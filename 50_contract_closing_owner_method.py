"""
Contract closing in the OWNER'S OWN STYLE (from his party sheets "DT Original Sheet", "June 25 - current",
"Wahab Fabric Closing"), built from the ERP (read only) into a Google Sheet so it can be compared with the
first closing file (48) and the set-length file (49).

Owner's rules (26-Sep-2026):
  - weight per meter = ends/inch x width x 1.0936 / 800 / closing count  (800, not 840; closing count = count - margin)
  - ends/inch: 46" qualities use the actual ends of the sized set (incl. selvedge, e.g. 12,520 / (3 x 46) = 90.7);
    wide qualities (118", 122" ...) use the quality ends (e.g. 76)
  - yarn consumed = weight per meter x NET FABRIC DISPATCHED (fabric OGP - fabric IGP). Fabric still in the mill
    is our internal record, not the party's - except where the party has not lifted fresh fabric (Wahab), where
    the owner adds the stock.
  - per yarn: bags received (1 bag = 100 lb) - bags consumed = balance; balance lb x rate per lb = amount
    (+ we hold party yarn = payable to party, - = receivable from party)
  - warp / weft split for display: warp bags = sets sized x (1 - normal shrinkage 10%) x warp lb/m, rest is weft
  - sets sized are listed in the corner to see the expected length
  - no sizing weights are used
"""

import importlib.util
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
_s = importlib.util.spec_from_file_location("cc", HERE / "48_contract_closing.py")
cc = importlib.util.module_from_spec(_s); _s.loader.exec_module(cc)
_s = importlib.util.spec_from_file_location("bs", HERE / "49_contract_closing_by_sets.py")
bs = importlib.util.module_from_spec(_s); _s.loader.exec_module(bs)
dash = cc.dash
KG_PER_LB = cc.KG_PER_LB
OUT = HERE / "financial_statements" / "dost_contract_closing_owner_method.xlsx"

NAVY, INPUT, SEC, GREEN, AMBER = cc.NAVY, cc.INPUT, cc.SEC, cc.GREEN, cc.AMBER
WHITE, B = cc.WHITE, cc.B
NUM0, NUM1, NUM4, PCT = cc.NUM0, cc.NUM1, cc.NUM4, cc.PCT
NUM2 = '#,##0.00;[Red](#,##0.00);"–"'
TITLE = Font(bold=True, size=14, color="1F3B57")
SUB = Font(bold=True, color="1F3B57")

NARROW_WIDTH = 60                  # qualities up to this width use the actual set ends
INCLUDE_STOCK = {43, 46}           # Wahab: fresh fabric not lifted, owner adds fabric stock
WARP_ALIAS = {46: {102, 120}}      # Wahab 46 warp came as two ERP codes (40/1 staple, 42/1 staple)
SAME_YARN = {6, 18}                # 52/1 CTN and 54/1 CTN: same yarn, same bags (owner)

# yarn OGP with a price in the description ("B#32 12/1 CTN 3 BAGS @29000/-") = mill yarn SOLD to the party, not a return
TT = "CASE WHEN m.tt = 'OGP_YARN' AND REGEXP_LIKE(m.tdesc, '@ *[0-9]') THEN 'OGP_SALE' ELSE m.tt END"
YCODE_SQL = f"""
 SELECT d.cont_id, {TT}, d.ycode, SUM({cc.KG})
 FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no = m.trans_no LEFT JOIN dostex.yarn y ON y.ycode = d.ycode
 WHERE m.tt IN ('IGP_YARN','OGP_YARN','OPN_SIZED_BEAMS','TO_SIZING','FROM_SIZING','YARN_TO_LOOMS','YARN_FROM_LOOMS')
 GROUP BY d.cont_id, {TT}, d.ycode"""
YDET_SQL = f"""
 SELECT d.cont_id, {TT}, m.tno, m.gp_no, m.td, y.yarn_desc, d.ycode, d.bags, d.cones, {cc.KG}, d.yarn_wt, m.tdesc
 FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no = m.trans_no LEFT JOIN dostex.yarn y ON y.ycode = d.ycode
 WHERE m.tt IN ('IGP_YARN','OGP_YARN','OPN_SIZED_BEAMS') ORDER BY m.td, m.tno"""


def hdr(ws, row, labels, col0=1, height=32):
    for i, h in enumerate(labels, col0):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = WHITE, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = height


def party_token(party):
    for (_, p) in list(cc.OVERRIDES) + list(cc.PARTY_RATES):
        if p in (party or ""):
            return p
    return ""


def main():
    cur = dash.connect(dash.load_config()).cursor()
    rows = dash.rows
    contracts = sorted(rows(cur, cc.CONTRACTS_SQL), key=lambda c: int(c[1]))
    ynames = {int(y): (d or "").strip() for y, d in rows(cur, cc.YARN_NAMES_SQL)}
    # 52/1 CTN (code 6) and 54/1 CTN (code 18) are the same yarn in the same bags (owner, 28-Sep-2026): book both
    # to whichever of the two the contract uses, so its counts (e.g. Azam Puri 54/1 = 51.5/52.5) still apply
    uses = {c[0]: {c[12], c[15]} | {y for y, *_ in cc.WEFT_MIX.get(int(c[1]), [])} for c in contracts}
    ymov = {}
    for cid, tt, yc, kg in rows(cur, YCODE_SQL):
        if yc is not None:
            yc = int(yc)
            if yc in SAME_YARN:
                yc = 18 if 18 in uses.get(cid, set()) else 6
            ymov[(cid, tt, yc)] = ymov.get((cid, tt, yc), 0.0) + float(kg or 0)
    ydet = {}
    for r in rows(cur, YDET_SQL):
        ydet.setdefault(r[0], []).append(r[1:])
    fdet = {}
    for r in rows(cur, cc.FAB_DETAIL_SQL):
        fdet.setdefault(r[0], []).append(r[1:])
    stock = {cid: float(v or 0) for cid, v in rows(cur, cc.STOCK_SQL)}
    bal = {an: float(v or 0) for an, v in rows(cur, cc.PARTY_BAL_SQL)}
    prod = {}
    for t, s, cid, m, f, l in rows(cur, bs.PROD_SQL):
        prod[(t, s)] = prod.get((t, s), 0) + float(m or 0)
    sets = {}
    for tr, sr, setno, bcode, blen, cid, tt, tno, gp, td, desc, ends, slen, sizer in rows(cur, bs.BEAMS_SQL):
        k = (cid, tr, setno)
        s = sets.setdefault(k, dict(cid=cid, setno=setno, tno=tno, td=td, sizer=sizer, tt=tt, ends=0, slen=0,
                                    beams=0, woven=0.0))
        s["ends"] = max(s["ends"], int(ends or 0))
        s["slen"] = max(s["slen"], float(slen or 0))
        s["beams"] += 1
        s["woven"] += prod.get((tr, sr), 0)

    wb = Workbook()
    # ---------------- Inputs
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = "Closing counts and rates (owner's) - edit the yellow cells, every contract recalculates"
    inp["A1"].font = TITLE
    inp["A2"] = "Normal set shrinkage (sets sized -> expected fabric):"
    inp["H2"] = 0.10
    inp["H2"].fill, inp["H2"].number_format = INPUT, "0.0%"
    cc.header(inp, 5, ["Key", "Yarn", "Party (blank = all)", "ERP yarn code", "Warp closing count", "Weft closing count",
                       "Rate Rs per lb", "Note"], [14, 26, 16, 10, 12, 12, 12, 40])
    party_entries = {}
    for (yc, p), (w, f) in cc.OVERRIDES.items():
        party_entries[(yc, p)] = [w, f, cc.YARNS[yc][3]]
    for (yc, p), rate in cc.PARTY_RATES.items():
        base = cc.YARNS.get(yc, ("", None, None, None))
        e = party_entries.setdefault((yc, p), [base[1], base[2], rate])
        e[2] = rate
    entries = [(yc, "", w, f, rate, "") for yc, (_, w, f, rate) in cc.YARNS.items() if isinstance(yc, int)]
    entries += [(yc, p, w, f, rate, "party-specific") for (yc, p), (w, f, rate) in sorted(party_entries.items(),
                                                                                         key=lambda x: str(x))]
    used = set()
    for c in contracts:
        used |= {c[12], c[15]}
    for (cid, tt, yc) in ymov:
        used.add(yc)
    entries += [(yc, "", None, None, None, "NEEDED - count and rate not given") for yc in sorted(
        {u for u in used if u is not None} - set(cc.YARNS))]
    r = 6
    for yc, p, w, f, rate, note in entries:
        for col, v in enumerate([f"{yc}|{p}", cc.YARNS[yc][0] if yc in cc.YARNS else ynames.get(int(yc), f"code {yc}"),
                                 p, yc, w, f, rate, note], 1):
            c = inp.cell(row=r, column=col, value=v)
            if col in (5, 6, 7):
                c.fill = INPUT
        r += 1
    last_inp = r - 1
    inp.freeze_panes = "A6"

    def key_for(yc, party):
        p = party_token(party)
        return f"{yc}|{p}" if p and (yc, p) in party_entries else f"{yc}|"

    def lk(col, key):
        return (f"IFERROR(INDEX(Inputs!${col}$6:${col}${last_inp},"
                f"MATCH(\"{key}\",Inputs!$A$6:$A${last_inp},0)),\"\")")

    # ---------------- Summary
    sm = wb.create_sheet("All contracts", 0)
    sm["A1"] = f"Contract closing - OWNER'S METHOD (yarn vs net fabric dispatched, / 800) - {date.today():%d-%b-%Y}"
    sm["A1"].font = TITLE
    sm["A2"] = ("+ = we hold the party's yarn (payable to party); - = party owes us yarn (receivable from party). "
                "ERP read only; closing in this sheet. Open each contract tab for the working.")
    sm["A2"].font = Font(italic=True, size=9)
    scols = ["Contract", "Party", "Status", "Quality", "Fabric OGP (m)", "Fabric IGP / returned (m)",
             "Net fabric dispatched (m)", "Fabric in mill stock (m)", "Stock included?", "Meters for closing",
             "Yarn received (bags)", "Yarn consumed (bags)", "Balance (bags)", "Yarn result (Rs)", "Result",
             "Expected fabric from sets (m)", "Still to come from sets (m)", "Flags"]
    cc.header(sm, 4, scols, [8, 22, 16, 30, 11, 11, 12, 11, 8, 12, 11, 11, 10, 13, 20, 12, 12, 70])
    srow = 5
    tab_names = set()
    for (cid, no, party, cdate, exp, cqty, f, l, mtr, ag, rj, q, wy, ends, cww, fy, picks, cfw, width, conv) in contracts:
        no = int(no)
        party = party or ""
        finished = l.date() <= cc.CUTOFF
        width = float(width or 0)
        reed_q = float(ends or 0)
        mix = cc.WEFT_MIX.get(no)
        picks_used = sum(p for *_, p in mix) if mix else float(picks or 0)
        csets = sorted([s for s in sets.values() if s["cid"] == cid], key=lambda s: (s["td"], str(s["setno"])))
        # beams across the loom for each set, and the actual ends/inch for 46" qualities
        for s in csets:
            s["mult"] = max(1, round(s["ends"] / (reed_q * width))) if s["ends"] and reed_q and width else 1
        wsum = sum(s["slen"] * s["mult"] for s in csets if s["ends"] and width)
        if width and width <= NARROW_WIDTH and wsum:
            reed = round(sum(s["ends"] / width * s["slen"] for s in csets if s["ends"]) / wsum, 2)
            reed_note = f"actual set ends / width (quality says {reed_q:g})"
        else:
            reed = reed_q
            reed_note = "quality ends (wide width)"

        # yarn per code
        def mv(tt, yc):
            return ymov.get((cid, tt, yc), 0.0)
        codes = {yc for (c_, tt, yc) in ymov if c_ == cid}
        warp_codes = ({wy} if wy is not None else set()) | WARP_ALIAS.get(no, set())
        weft_parts = [(y, cnt, p) for y, cnt, p in mix] if mix else ([(fy, None, picks_used)] if fy is not None else [])
        weft_codes = {y for y, _, _ in weft_parts}
        yrows = []   # (label, codes, role, weft_part or None)
        same = wy is not None and wy in weft_codes
        if same:
            wp = next(p for p in weft_parts if p[0] == wy)
            yrows.append(("Warp + weft", sorted(warp_codes), wp))
            for p in weft_parts:
                if p[0] != wy:
                    yrows.append(("Weft", [p[0]], p))
        else:
            if warp_codes:
                yrows.append(("Warp", sorted(warp_codes), None))
            for p in weft_parts:
                yrows.append(("Weft", [p[0]], p))
        listed = set().union(*[set(x[1]) for x in yrows]) if yrows else set()
        for yc in sorted(codes - listed):
            netk = mv("IGP_YARN", yc) - mv("OGP_YARN", yc) + mv("OPN_SIZED_BEAMS", yc)
            if abs(netk) > 1:
                yrows.append(("Not in contract", [yc], None))

        short = re.sub(r"[^A-Z0-9 ]", "", party.upper()).replace("TRADERSEMB", "").split(" ")[0] or "PARTY"
        name = f"{no} {short}"[:28]
        while name in tab_names:
            name += "."
        tab_names.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = "2E8B57" if finished else "E07B00"
        ws["A1"] = (f"Contract {no} - {party} - CLOSING (owner's method)" if finished else
                    f"Contract {no} - {party} - PENDING CLOSING (still running, last production {l:%d-%b-%Y})")
        ws["A1"].font = Font(bold=True, size=14, color="1F3B57" if finished else "9A3412")
        flags = []

        # fabric
        fab = fdet.get(cid, [])
        ogp = sum(float(x[9] or 0) for x in fab if x[0] == "GPOUT")
        igp = sum(float(x[9] or 0) for x in fab if x[0] == "GPIN")
        info = [("Quality", q), ("Ends / inch used", reed), ("Picks / inch", picks_used), ("Width (in)", width),
                ("Fabric OGP - dispatched (m)", round(ogp)), ("Fabric IGP - returned (m)", round(igp)),
                ("Net fabric dispatched (m) = OGP - IGP", "=B8-B9"), ("Fabric in mill stock (m, ERP)", round(stock.get(cid, 0))),
                ("Include stock in closing? (Yes/No)", "Yes" if no in INCLUDE_STOCK else "No"),
                ("METERS FOR CLOSING", '=B10+IF(UPPER(B12)="YES",B11,0)'),
                ("Contract quantity (m)", cqty), ("Meters woven (ERP production)", round(float(mtr or 0)))]
        for i, (lab, v) in enumerate(info, 4):
            ws.cell(row=i, column=1, value=lab)
            c = ws.cell(row=i, column=2, value=v)
            if i in (8, 9, 10, 11, 13, 14, 15):
                c.number_format = NUM0
        for rr in (5, 6, 12):
            ws.cell(row=rr, column=2).fill = INPUT
        ws["C5"] = reed_note
        ws["C5"].font = Font(italic=True, size=9)
        if mix:
            ws["C6"] = "two wefts, pattern 11:1 (owner's Wahab closing)"
            ws["C6"].font = Font(italic=True, size=9)
        for rr in (10, 13):
            ws.cell(row=rr, column=1).font = B
            ws.cell(row=rr, column=2).font = B

        # yarn table
        yh = 17
        ws.cell(row=yh - 1, column=1, value="YARN ACCOUNT (1 bag = 100 lb)").font = SUB
        hdr(ws, yh, ["Yarn", "ERP code(s)", "Used as", "Warp closing count", "Weft closing count", "Weft picks",
                     "Warp lb / m", "Weft lb / m", "Total lb / m", "Yarn IGP (kg)", "Yarn OGP returned (kg)",
                     "Sent direct to sizer (enter) / opening beams (kg)", "Net received (kg)", "Net received (bags)",
                     "Consumed (bags)", "Yarn left in mill, goes back to party (bags)", "Balance (bags)",
                     "Balance (lb)", "Rate / lb", "Amount (Rs)", "Result"])
        r = yh + 1
        first = r
        warp_row = None
        for role, ycs, wp in yrows:
            label = " + ".join(cc.YARNS[y][0] if y in cc.YARNS else ynames.get(y, f"code {y}") for y in ycs)
            igk = sum(mv("IGP_YARN", y) for y in ycs)
            ogk = sum(mv("OGP_YARN", y) for y in ycs)
            opn = sum(mv("OPN_SIZED_BEAMS", y) for y in ycs)
            direct = 0.0
            if role in ("Warp", "Warp + weft"):
                tosz = sum(mv("TO_SIZING", y) - mv("FROM_SIZING", y) for y in ycs)
                if tosz - (igk - ogk) > 50:
                    flags.append(f"ERP shows {tosz - (igk - ogk):,.0f} kg more warp sent to sizing than yarn IGP on this "
                                 f"contract - if the party sent yarn direct to the sizer, enter it in the yellow cell")
            k0 = ycs[0]
            key = key_for(k0, party)
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=", ".join(str(y) for y in ycs))
            ws.cell(row=r, column=3, value=role)
            if role in ("Warp", "Warp + weft"):
                ws.cell(row=r, column=4, value="=" + lk("E", key))
                ws.cell(row=r, column=7, value=f"=IF(N(D{r})>0,$B$5*$B$7*1.0936/800/D{r},0)")
            else:
                ws.cell(row=r, column=7, value=0)
            if wp is not None:
                y, cnt, p = wp
                if cnt is None:
                    ws.cell(row=r, column=5, value="=" + lk("F", key_for(y, party)))
                    ws.cell(row=r, column=6, value="=$B$6")
                else:
                    ws.cell(row=r, column=5, value=cnt).fill = INPUT
                    ws.cell(row=r, column=6, value=round(p, 3)).fill = INPUT
                ws.cell(row=r, column=8, value=f"=IF(N(E{r})>0,F{r}*$B$7*1.0936/800/E{r},0)")
            else:
                ws.cell(row=r, column=8, value=0)
            ws.cell(row=r, column=9, value=f"=G{r}+H{r}")
            ws.cell(row=r, column=10, value=round(igk, 1))
            ws.cell(row=r, column=11, value=round(ogk, 1))
            ws.cell(row=r, column=12, value=round(direct + opn, 1))
            if role in ("Warp", "Warp + weft"):
                ws.cell(row=r, column=12).fill = INPUT
            ws.cell(row=r, column=13, value=f"=J{r}-K{r}+L{r}")
            ws.cell(row=r, column=14, value=f"=M{r}/{KG_PER_LB}/100")
            ws.cell(row=r, column=15, value=f"=I{r}*$B$13/100")
            ws.cell(row=r, column=16, value=0).fill = INPUT
            ws.cell(row=r, column=17, value=f"=N{r}-O{r}-P{r}")
            ws.cell(row=r, column=18, value=f"=Q{r}*100")
            ws.cell(row=r, column=19, value="=" + lk("G", key))
            ws.cell(row=r, column=20, value=f"=IFERROR(R{r}*S{r},0)")
            ws.cell(row=r, column=21, value=f'=IF(R{r}>=0,"Payable to party","Receivable from party")')
            for col in (4, 5, 6):
                ws.cell(row=r, column=col).number_format = "0.###"
            for col in (7, 8, 9):
                ws.cell(row=r, column=col).number_format = NUM4
            for col in (10, 11, 12, 13, 18, 20):
                ws.cell(row=r, column=col).number_format = NUM0
            for col in (14, 15, 16, 17):
                ws.cell(row=r, column=col).number_format = NUM2
            if role in ("Warp", "Warp + weft") and warp_row is None:
                warp_row = r
            if role == "Not in contract":
                flags.append(f"{label}: yarn booked to this contract but not in its construction - check")
                ws.cell(row=r, column=3).fill = AMBER
            if all(y not in cc.YARNS for y in ycs):
                flags.append(f"{label}: count/rate not given")
            r += 1
        last = r - 1
        tot = r
        ws.cell(row=tot, column=1, value="TOTAL").font = B
        for col in (10, 11, 12, 13, 14, 15, 16, 17, 18, 20):
            L = get_column_letter(col)
            c = ws.cell(row=tot, column=col, value=f"=SUM({L}{first}:{L}{last})" if last >= first else 0)
            c.font = B
            c.number_format = NUM2 if col in (14, 15, 16, 17) else NUM0
        res = tot + 1
        ws.cell(row=res, column=1, value="YARN RESULT").font = Font(bold=True, size=12, color="1F3B57")
        ws.cell(row=res, column=20, value=f"=T{tot}").number_format = NUM0
        ws.cell(row=res, column=20).font = Font(bold=True, size=12)
        ws.cell(row=res, column=21, value=f'=IF(T{tot}>=0,"Payable to {party}","Receivable from {party}")').font = B

        # owner-style warp / weft split
        sp = res + 2
        ws.cell(row=sp, column=1, value="WARP / WEFT SPLIT (owner style: warp bags = sets sized x (1 - shrinkage) x warp lb/m)").font = SUB
        if warp_row:
            lines = [("Expected fabric from sets sized (m)", "=EXPFAB"),
                     ("Warp bags for the sets sized", f"=B{sp + 1}*G{warp_row}/100"),
                     ("Warp bags consumed (meters for closing)", f"=$B$13*G{warp_row}/100"),
                     ("Warp balance (bags)", f"=B{sp + 2}-B{sp + 3}")]
            if same:
                lines += [("Weft bags = received - warp bags for sets", f"=N{warp_row}-B{sp + 2}"),
                          ("Weft bags consumed", f"=$B$13*H{warp_row}/100"),
                          ("Weft balance (bags)", f"=B{sp + 5}-B{sp + 6}")]
            else:
                lines += [("Warp bags actually received", f"=N{warp_row}"),
                          ("Received - needed for sets (bags)", f"=B{sp + 5}-B{sp + 2}")]
            for i, (lab, v) in enumerate(lines, 1):
                ws.cell(row=sp + i, column=1, value=lab)
                ws.cell(row=sp + i, column=2, value=v).number_format = NUM0 if i == 1 else NUM2
            split_end = sp + len(lines)
        else:
            split_end = sp
        # sets sized (right corner)
        sc = 23
        ws.cell(row=3, column=sc, value="SETS SIZED (expected length)").font = SUB
        hdr(ws, 4, ["Set #", "Received", "Sizer", "Ends", "Set length per beam (m)", "Beams across",
                    "Set length (m)", "Shrinkage %", "Expected fabric (m)", "Fabric woven from set (ERP m)",
                    "Actual shrinkage %"], col0=sc)
        rr = 5
        s_first = rr
        for s in csets:
            vals = [s["setno"], s["td"], s["sizer"], s["ends"], s["slen"], s["mult"]]
            for i, v in enumerate(vals):
                c = ws.cell(row=rr, column=sc + i, value=v)
                if i == 1:
                    c.number_format = "dd-mmm-yy"
                elif i in (3, 4):
                    c.number_format = NUM0
            L = lambda k: get_column_letter(sc + k)  # noqa: E731
            ws.cell(row=rr, column=sc + 6, value=f"={L(4)}{rr}*{L(5)}{rr}").number_format = NUM0
            ws.cell(row=rr, column=sc + 7, value="=Inputs!$H$2").number_format = PCT
            ws.cell(row=rr, column=sc + 7).fill = INPUT
            ws.cell(row=rr, column=sc + 8, value=f"={L(6)}{rr}*(1-{L(7)}{rr})").number_format = NUM0
            ws.cell(row=rr, column=sc + 9, value=round(s["woven"])).number_format = NUM0
            ws.cell(row=rr, column=sc + 10, value=f"=IF({L(6)}{rr}>0,1-{L(9)}{rr}/{L(6)}{rr},\"\")").number_format = PCT
            if s["tt"] == "OPN_SIZED_BEAMS":
                ws.cell(row=rr, column=sc).fill = AMBER
            rr += 1
        s_last = rr - 1
        L = lambda k: get_column_letter(sc + k)  # noqa: E731
        ws.cell(row=rr, column=sc, value="TOTAL").font = B
        for k in (6, 8, 9):
            c = ws.cell(row=rr, column=sc + k, value=f"=SUM({L(k)}{s_first}:{L(k)}{s_last})" if s_last >= s_first else 0)
            c.font, c.number_format = B, NUM0
        s_tot = rr
        exp_ref = f"{L(8)}{s_tot}"
        for i, (lab, v) in enumerate([("Expected fabric from sets", f"={exp_ref}"), ("Net fabric dispatched", "=B10"),
                                      ("Fabric in mill stock", "=B11"),
                                      ("Still to come from the sets (m)", f"={exp_ref}-B10-B11")], 2):
            ws.cell(row=s_tot + i, column=sc, value=lab)
            ws.cell(row=s_tot + i, column=sc + 8, value=v).number_format = NUM0
        ws.cell(row=s_tot + 5, column=sc).font = B
        if warp_row:
            ws.cell(row=sp + 1, column=2, value=f"={exp_ref}")
        if not csets:
            flags.append("No sized sets on this contract in the ERP")

        # gate passes
        g = max(split_end, 0) + 3
        ws.cell(row=g, column=1, value="YARN GATE PASSES").font = SUB
        hdr(ws, g + 1, ["Date", "IGP / OGP #", "Type", "Yarn", "Bags", "Cones", "Kg used", "Bags of 100 lb",
                        "Kg typed in ERP", "Description"])
        rr = g + 2
        sold = 0.0
        for tt, tno, gp, td, yd, yc, bg, cn, kg, typed, desc in ydet.get(cid, []):
            sign = {"OGP_YARN": -1, "OGP_SALE": 0}.get(tt, 1)
            if tt == "OGP_SALE":
                sold += float(kg or 0)
            vals = [td, gp or tno, {"IGP_YARN": "IGP received", "OGP_YARN": "OGP returned",
                                    "OGP_SALE": "OGP - mill yarn SOLD to party (not a return)",
                                    "OPN_SIZED_BEAMS": "Opening sized beams"}.get(tt, tt), yd,
                    bg, cn, (sign or 1) * float(kg or 0), (sign or 1) * float(kg or 0) / KG_PER_LB / 100, typed, desc]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=rr, column=i, value=v)
                c.number_format = {1: "dd-mmm-yy", 7: NUM0, 8: NUM2, 9: NUM1}.get(i, "General")
            if typed is not None and abs(float(typed or 0) - float(kg or 0)) > 1:
                ws.cell(row=rr, column=9).fill = AMBER
            if tt == "OGP_SALE":
                ws.cell(row=rr, column=3).fill = AMBER
            rr += 1
        if sold:
            flags.append(f"{sold:,.0f} kg mill yarn sold to the party on priced OGPs - not treated as a return")
        if any(t is not None and abs(float(t or 0) - float(k or 0)) > 1
               for tt, *_, k, t, _d in ydet.get(cid, []) if tt == "IGP_YARN"):
            flags.append("Some yarn IGP weights in the ERP are wrong (0 / few kg / whole gate-pass total) - bags x 100 lb used")
        rr += 1
        ws.cell(row=rr, column=1, value="FABRIC GATE PASSES (net dispatched = OGP - IGP)").font = SUB
        hdr(ws, rr + 1, ["Date", "OGP / IGP #", "Type", "Meters", "A grade", "B + C + CP", "Delivered to", "Invoice #",
                         "Rate"])
        rr += 2
        for tt, gp, trn, td, inv, pty, dto, a, rjm, tm, rate, amt in fab:
            sign = -1 if tt == "GPIN" else 1
            vals = [td, gp or trn, "OGP dispatched" if tt == "GPOUT" else "IGP returned", sign * float(tm or 0),
                    a, rjm, dto, inv, rate]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=rr, column=i, value=v)
                c.number_format = {1: "dd-mmm-yy", 4: NUM0, 5: NUM0, 6: NUM0}.get(i, "General")
            rr += 1

        if cdate and cdate.date() <= date(2023, 10, 1):
            flags.append("Started before the new ERP: old-ERP yarn/fabric has no contract number - incomplete")
        if no in INCLUDE_STOCK:
            flags.append("Fabric stock included (fresh fabric not lifted by the party)")
        if mix:
            flags.append("Two wefts: 11:1 pick pattern from the owner's Wahab closing")
        ws.cell(row=2, column=1, value="; ".join(flags)).font = Font(italic=True, size=9, color="9A3412")
        for col, w in zip(range(1, 22), (34, 12, 13, 9, 9, 8, 9, 9, 9, 11, 11, 13, 11, 10, 10, 12, 10, 10, 8, 12, 20)):
            ws.column_dimensions[get_column_letter(col)].width = w
        for col, w in zip(range(sc, sc + 11), (8, 10, 14, 8, 10, 7, 10, 9, 11, 11, 9)):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.column_dimensions["V"].width = 3
        ws.freeze_panes = "A4"

        # summary row
        T = f"'{name}'!"
        vals = [no, party, "Finished" if finished else "PENDING CLOSING", q, f"={T}B8", f"={T}B9", f"={T}B10",
                f"={T}B11", f"={T}B12", f"={T}B13", f"={T}N{tot}", f"={T}O{tot}", f"={T}Q{tot}", f"={T}T{tot}",
                f'=IF(N{srow}>=0,"Payable to party","Receivable from party")', f"={T}{exp_ref}",
                f"={T}{exp_ref}-{T}B10-{T}B11", "; ".join(flags)]
        for i, v in enumerate(vals, 1):
            c = sm.cell(row=srow, column=i, value=v)
            c.number_format = {11: NUM2, 12: NUM2, 13: NUM2}.get(i, NUM0 if 5 <= i <= 17 and i != 9 else "General")
        sm.cell(row=srow, column=3).fill = GREEN if finished else AMBER
        srow += 1
    sm.freeze_panes = "C5"
    sm.auto_filter.ref = f"A4:R{srow - 1}"
    last_sm = srow - 1

    # ---------------- Party summary (finished contracts) with the ledger balance
    ps = wb.create_sheet("Party final hisaab", 1)
    ps["A1"] = "Final hisaab by party - yarn result of FINISHED contracts + conversion ledger balance"
    ps["A1"].font = TITLE
    ps["A2"] = ("Yarn result + = we hold party yarn (payable); ledger balance + = party owes us (receivable). "
                "Net receivable = ledger balance - yarn result.")
    ps["A2"].font = Font(italic=True, size=9)
    cc.header(ps, 4, ["Party", "Finished contracts", "Yarn result - finished (Rs)", "Yarn result - pending (Rs, info)",
                      "Conversion ledger balance today (Rs)", "NET RECEIVABLE from party (Rs)", "Result"],
              [30, 12, 16, 16, 18, 18, 22])
    parties = sorted({(c[2] or "") for c in contracts})
    rng = lambda col: f"'All contracts'!${col}$5:${col}${last_sm}"  # noqa: E731
    r = 5
    for p in parties:
        ps.cell(row=r, column=1, value=p)
        ps.cell(row=r, column=2, value=f"=COUNTIFS({rng('B')},$A{r},{rng('C')},\"Finished\")")
        ps.cell(row=r, column=3, value=f"=SUMIFS({rng('N')},{rng('B')},$A{r},{rng('C')},\"Finished\")").number_format = NUM0
        ps.cell(row=r, column=4, value=f"=SUMIFS({rng('N')},{rng('B')},$A{r},{rng('C')},\"PENDING CLOSING\")").number_format = NUM0
        ps.cell(row=r, column=5, value=round(bal.get(p, 0))).number_format = NUM0
        ps.cell(row=r, column=6, value=f"=E{r}-C{r}").number_format = NUM0
        ps.cell(row=r, column=7, value=f'=IF(F{r}>=0,"Receivable from party","Payable to party")')
        r += 1
    ps.freeze_panes = "B5"

    # ---------------- Check against the owner's own Wahab closing
    ck = wb.create_sheet("Check vs owner's Wahab", 2)
    ck["A1"] = "Check: this method vs the owner's own 'Wahab Fabric Closing' (weights per meter and bags at his meters)"
    ck["A1"].font = TITLE
    cc.header(ck, 3, ["Contract", "Yarn", "Owner lb/m", "This sheet lb/m", "Diff %", "Owner meters",
                      "Owner bags consumed", "This sheet bags at owner's meters"], [9, 28, 11, 12, 9, 12, 14, 16])
    rows_ck = [(43, "60/1 CTN warp", 0.0984, 90.76, None, 58, 40000, None),
               (43, "60/1 CTN weft", 0.0625, None, 64 * 11 / 12, 59, 40000, None),
               (43, "12/1 CTN weft", 0.0305, None, 64 / 12, 11, 40000, None),
               (46, "40 staple warp", 0.1426, 90.72, None, 40, 25801, None),
               (46, "120/1 VISCOSE weft", 0.0878, None, 64 * 11 / 12, 42, 25801, None),
               (46, "20/2 VISCOSE (10s) weft", 0.0373, None, 64 / 12, 9, 25801, None)]
    owner_bags = {("60/1 CTN warp"): 39.3445, ("60/1 CTN weft"): 25.0107, ("12/1 CTN weft"): 12.1953,
                  ("40 staple warp"): 36.7983, ("120/1 VISCOSE weft"): 22.6624, ("20/2 VISCOSE (10s) weft"): 9.6143}
    r = 4
    for no_, yl, own, reed_, pk, cnt, mtr_, _ in rows_ck:
        ours = (reed_ if reed_ else pk) * 46 * 1.0936 / 800 / cnt
        vals = [no_, yl, own, round(ours, 4), ours / own - 1, mtr_, owner_bags[yl], ours * mtr_ / 100]
        for i, v in enumerate(vals, 1):
            c = ck.cell(row=r, column=i, value=v)
            c.number_format = {3: NUM4, 4: NUM4, 5: PCT, 6: NUM0, 7: NUM2, 8: NUM2}.get(i, "General")
        r += 1
    ck.cell(row=r + 1, column=1, value=("Ends/inch from the sized sets: contract 43 = 12,520 / (3 x 46) = 90.7; contract 46 = "
                                        "12,519 / (3 x 46) = 90.7. Picks 64 split 11:1. Owner's meters: 43 = 39,027 dispatched "
                                        "+ 973 stock; 46 = 17,751 dispatched + 8,050 stock (at the time of his closing)."))

    # ---------------- How this works
    hw = wb.create_sheet("How this works")
    for i, t in enumerate([
        "OWNER'S METHOD - read from the owner's own party sheets (KB EXPORTS, PC, Mehak, Khaddar, Viscose, Wahab closing)",
        "Weight per meter (lb) = ends per inch x width x 1.0936 / 800 / closing count (warp) and picks x width x 1.0936 / 800 / closing count (weft).",
        "800 instead of 840 is the owner's built-in allowance (about 5%). Closing count = count - margin (e.g. 54s: 51.5 warp / 52.5 weft).",
        "Ends per inch: 46\" qualities use the actual ends of the sized set (incl. selvedge); wide qualities use the quality ends.",
        "Meters for closing = net fabric dispatched (fabric OGP - fabric IGP). Stock in the mill is our internal record, "
        "not the party's - except when the party has not lifted fresh fabric (Wahab: stock added, switch on each tab).",
        "Yarn received = yarn IGP - yarn OGP (returns), in kg and bags of 100 lb. Where the typed kg is not "
        "believable (under 15 or over 60 kg per bag) bags x 100 lb is used. Priced yarn OGPs (\"@29000/-\") are mill yarn "
        "SOLD to the party, not returns. Yarn the party sent direct to the sizer can be entered in the yellow cell.",
        "Balance = bags received - bags consumed - yarn left in the mill that goes back to the party (yellow, fill in).",
        "Amount = balance lb x owner's rate per lb. + = we hold the party's yarn (payable to party); - = receivable from party.",
        "Warp/weft split (owner style): warp bags = sets sized x (1 - normal shrinkage) x warp lb/m; the rest is weft.",
        "Sets sized are on the right of each tab: set length x beams across x (1 - shrinkage) = expected fabric.",
        "Party final hisaab = conversion ledger balance - yarn result of finished contracts.",
        "Contracts started before Oct-2023 are incomplete (old ERP yarn and fabric have no contract number).",
        "ERP data is only read; nothing is written back to Oracle.",
    ], 1):
        hw.cell(row=i, column=1, value=t).font = TITLE if i == 1 else Font()
    hw.column_dimensions["A"].width = 150

    wb.move_sheet("All contracts", -(len(wb.sheetnames)))
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"contracts: {len(contracts)} tabs: {len(wb.sheetnames)}")


if __name__ == "__main__":
    main()
    if len(sys.argv) > 2 and sys.argv[1] == "--upload":
        cc.OUT = OUT
        cc.upload(sys.argv[2])
