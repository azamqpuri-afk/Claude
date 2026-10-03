"""
58_pilferage_check_pc.py - day-by-day check of Haroon Fabric's 30/1 PC yarn (ERP code 109, "31 PC") for signs of
pilferage, same checks as 54_pilferage_check.py (Qadri 52/54). Read only.
 1. Weft issued to looms vs weft needed for the fabric woven (7-day windows, since production is entered at roll check)
 2. Entry pattern: issues with no production nearby, very large issues, entries keyed long after the date or out of hours
 3. Sizing: yarn sent vs sized-beam weight per sizer and contract
 4. Per contract: weft issued vs weft needed, and 30/1 PC issued to looms under non-PC contracts
Writes an Excel workbook to the Drive audit folder.
"""
import importlib.util
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
_s = importlib.util.spec_from_file_location("cc", HERE / "48_contract_closing.py")
cc = importlib.util.module_from_spec(_s); _s.loader.exec_module(cc)
rows = cc.dash.rows
YCODE = 109
WARP_COUNT, WEFT_COUNT = 29, 30   # owner: "31 PC" closes at warp 29, weft 30
FROM = date(2023, 10, 1)
RED = PatternFill("solid", fgColor="F8C9C9"); AMB = PatternFill("solid", fgColor="FFE3B3")


def main():
    cur = cc.dash.connect(cc.dash.load_config()).cursor()
    cons = {}; warp_m = {}                                 # cont_no -> weft / warp kg per fabric meter
    for r in rows(cur, cc.CONTRACTS_SQL):
        if r[15] == YCODE:
            cons[int(r[1])] = float(r[16]) * float(r[18]) * 1.0936 / 800 / WEFT_COUNT * 0.45359237
            warp_m[int(r[1])] = float(r[13]) * float(r[18]) * 1.0936 / 800 / WARP_COUNT * 0.45359237
    cl = ",".join(str(c) for c in sorted(cons))
    issues = defaultdict(float); ilines = []
    for d, tno, cno, bags, cones, kg, feed, user in rows(cur, f"""
        SELECT TRUNC(m.td), m.tno, c.cont_no, NVL(d.bags,0), NVL(d.cones,0), {cc.KG}, m.feedd, m.usern
        FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no LEFT JOIN dostex.yarn y ON y.ycode=d.ycode
        JOIN dostex.contm c ON c.cont_id=d.cont_id
        WHERE m.tt='YARN_TO_LOOMS' AND d.ycode = {YCODE} AND c.cont_no IN ({cl}) AND m.td >= :f0""", f0=FROM):
        d = d.date(); kg = float(kg or 0)
        issues[d] += kg
        ilines.append([d, tno, int(cno), float(bags), float(cones), round(kg, 1), feed, (user or "").strip()])
    back = defaultdict(float)
    for d, kg in rows(cur, f"""SELECT TRUNC(m.td), SUM({cc.KG}) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id
        WHERE m.tt='YARN_FROM_LOOMS' AND d.ycode = {YCODE} AND c.cont_no IN ({cl}) AND m.td >= :f0 GROUP BY TRUNC(m.td)""", f0=FROM):
        back[d.date()] += float(kg or 0)
    prod = defaultdict(float); need = defaultdict(float); c_prod = defaultdict(float)
    for d, cno, m in rows(cur, f"""SELECT TRUNC(td), cont_no, SUM(total_mtr) FROM dostex.wv_prod_v WHERE cont_no IN ({cl}) AND td >= :f0
                                  GROUP BY TRUNC(td), cont_no""", f0=FROM):
        d = d.date(); m = float(m or 0)
        prod[d] += m; need[d] += m * cons[int(cno)]; c_prod[int(cno)] += m
    c_last = {int(n): d for n, d in rows(cur, f"SELECT cont_no, TO_CHAR(MAX(td),'YYYY-MM-DD') FROM dostex.wv_prod_v WHERE cont_no IN ({cl}) GROUP BY cont_no")}
    c_warp = defaultdict(float)
    for cno, tt, kg in rows(cur, f"""SELECT c.cont_no, m.tt, SUM({cc.KG}) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id
        WHERE m.tt IN ('TO_SIZING','FROM_SIZING') AND d.ycode = {YCODE} AND c.cont_no IN ({cl}) GROUP BY c.cont_no, m.tt"""):
        c_warp[int(cno)] += float(kg or 0) * (1 if tt == "TO_SIZING" else -1)
    c_iss = defaultdict(float)
    for l in ilines:
        c_iss[l[2]] += l[5]
    other = rows(cur, f"""SELECT c.cont_no, MAX(a.an), ROUND(SUM({cc.KG})) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id LEFT JOIN dostex.ach a ON a.ac=c.party_code
        WHERE m.tt='YARN_TO_LOOMS' AND d.ycode = {YCODE} AND c.cont_no NOT IN ({cl}) GROUP BY c.cont_no""")

    days = []
    d, end = FROM, date.today()
    while d <= end:
        days.append(d); d += timedelta(days=1)
    daily = []
    for i, d in enumerate(days):
        w = days[max(0, i - 6): i + 1]
        iss = sum(issues[x] - back[x] for x in w); nd = sum(need[x] for x in w)
        daily.append([d, round(issues[d] - back[d], 1), round(prod[d]), round(need[d], 1), round(iss), round(nd), round(iss - nd)])
    by_day = {r[0]: r for r in daily}

    flags = []
    by_day_lines = defaultdict(list)
    for l in ilines:
        by_day_lines[l[0]].append(l)
    for d, L in by_day_lines.items():
        tot = sum(x[5] for x in L)
        near = sum(prod[d + timedelta(k)] for k in range(-3, 4))
        wk = by_day[d]
        reasons = []
        if tot >= 300 and near == 0:
            reasons.append(f"{tot:,.0f} kg issued but no production entered within 3 days either side")
        if tot >= 700:
            reasons.append(f"very large day: {tot:,.0f} kg (≈{tot / 45.36:.0f} bags) issued")
        if wk[6] > 400 and wk[4] > 1.3 * max(wk[5], 1):
            reasons.append(f"week to this day: {wk[4]:,} kg issued vs {wk[5]:,} kg needed for fabric woven (+{wk[6]:,} kg)")
        for x in L:
            feed, user = x[6], x[7]
            if feed and (feed.date() - d).days >= 7:
                reasons.append(f"issue {x[1]} keyed in {(feed.date() - d).days} days after its date (on {feed:%d-%b-%y} by {user or '?'})")
            if feed and (feed.hour >= 21 or feed.hour < 6):
                reasons.append(f"issue {x[1]} keyed at {feed:%H:%M} (outside working hours) by {user or '?'}")
        if reasons:
            flags.append([d, round(tot, 1), round(prod[d]), "; ".join(dict.fromkeys(reasons)), len(reasons)])

    sent = defaultdict(float)
    for an, cno, tt, kg in rows(cur, f"""SELECT a.an, c.cont_no, m.tt, SUM({cc.KG}) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id LEFT JOIN dostex.ach a ON a.ac=NVL(m.ac1,m.ac2)
        WHERE m.tt IN ('TO_SIZING','FROM_SIZING') AND d.ycode = {YCODE} AND m.td >= :f0 GROUP BY a.an, c.cont_no, m.tt""", f0=FROM):
        sent[((an or "").strip(), int(cno))] += float(kg or 0) * (1 if tt == "TO_SIZING" else -1)
    rec = defaultdict(float); lastd = {}
    for d, an, cno, desc, kg in rows(cur, f"""SELECT TRUNC(m.td), a.an, c.cont_no, m.tdesc, SUM({cc.KG}) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id LEFT JOIN dostex.ach a ON a.ac=NVL(m.ac1,m.ac2)
        WHERE m.tt='IGP_SIZED_BEAMS' AND d.ycode = {YCODE} AND m.td >= :f0 GROUP BY TRUNC(m.td), a.an, c.cont_no, m.tdesc""", f0=FROM):
        k = ((an or "").strip(), int(cno))
        actual = cc.desc_kg(desc)
        rec[k] += actual if actual else float(kg or 0)
        lastd[k] = max(lastd.get(k, d.date()), d.date())
    sizing = []
    for k, s in sent.items():
        r = rec.get(k, 0)
        if s > 0 and r > 0:
            sizing.append([k[0], k[1], round(s), round(r), round(s - r), round((s - r) / s * 100, 1), lastd.get(k)])

    wb = Workbook(); ws = wb.active; ws.title = "Days to check"
    ws.append(["Date", "Weft issued that day (kg)", "Fabric woven that day (m)", "Why flagged", "Signals"])
    for r in sorted(flags, key=lambda r: (-r[4], -r[1])):
        ws.append(r)
        ws.cell(row=ws.max_row, column=1).fill = RED if r[4] >= 2 else AMB
    ws2 = wb.create_sheet("Daily weft vs production")
    ws2.append(["Date", "Weft issued net (kg)", "Fabric woven (m)", "Weft needed for it (kg)", "Issued, last 7 days (kg)",
                "Needed, last 7 days (kg)", "Excess, last 7 days (kg)"])
    for r in daily:
        ws2.append(r)
        if r[6] > 400 and r[4] > 1.3 * max(r[5], 1):
            ws2.cell(row=ws2.max_row, column=7).fill = RED
    ws5 = wb.create_sheet("By contract")
    ws5.append(["Contract", "Fabric woven (m)", "Weft kg / m (count 30)", "Weft needed (kg)", "Weft issued (kg)", "Weft difference (kg)", "Weft difference %",
                "Warp kg / m (count 29)", "Warp needed for fabric woven (kg)", "Warp sent to sizing, net (kg)", "Warp difference (kg)", "Warp difference %", "Last woven"])
    for cno in sorted(cons):
        nd = c_prod[cno] * cons[cno]; iss = c_iss[cno]
        wn = c_prod[cno] * warp_m[cno]; ws_ = c_warp[cno]
        ws5.append([cno, round(c_prod[cno]), round(cons[cno], 5), round(nd), round(iss), round(iss - nd), round((iss - nd) / nd * 100, 1) if nd else None,
                    round(warp_m[cno], 5), round(wn), round(ws_), round(ws_ - wn), round((ws_ - wn) / wn * 100, 1) if wn else None, c_last.get(cno)])
    ws5.append(["TOTAL", round(sum(c_prod.values())), "", round(sum(c_prod[c] * cons[c] for c in cons)), round(sum(c_iss.values())), "", "",
                "", round(sum(c_prod[c] * warp_m[c] for c in cons)), round(sum(c_warp[c] for c in cons))])
    ws5.append(["Running contracts still have warp on looms / beams, so their warp difference is not final."])
    ws5.append([]); ws5.append(["30/1 PC issued to looms under other contracts (kg):"])
    for r in other:
        ws5.append(list(r))
    ws3 = wb.create_sheet("Sizing loss by sizer")
    ws3.append(["Sizer", "Contract", "Yarn sent net (kg)", "Sized beams back, actual kg", "Difference (kg)", "Loss %", "Last beams back"])
    for r in sorted(sizing, key=lambda r: -r[5]):
        ws3.append(r)
        if r[5] > 6:
            ws3.cell(row=ws3.max_row, column=6).fill = RED
    ws4 = wb.create_sheet("All weft issues")
    ws4.append(["Date", "Issue no.", "Contract", "Bags", "Cones", "Kg", "Keyed at", "User"])
    for l in sorted(ilines):
        ws4.append(l)
    for s in wb.worksheets:
        for c in s[1]:
            c.font = Font(bold=True)
    out = Path(cc.dash.load_config()["output"]["folder"]) / "Claude ERP Audit" / f"Haroon 31 PC yarn day-by-day check {date.today():%Y-%m-%d}.xlsx"
    wb.save(out)
    tot_iss = sum(r[1] for r in daily); tot_need = sum(r[3] for r in daily)
    print(f"contracts {cl}")
    print(f"period {FROM} to {date.today()}: weft issued {tot_iss:,.0f} kg vs needed {tot_need:,.0f} kg ({(tot_iss - tot_need) / tot_need * 100:+.1f}%)")
    print("BY CONTRACT:")
    for cno in sorted(cons):
        nd = c_prod[cno] * cons[cno]
        wn = c_prod[cno] * warp_m[cno]
        print(cno, c_last.get(cno), round(c_prod[cno]), "weft", round(nd), round(c_iss[cno]), f"{(c_iss[cno] - nd) / nd * 100:+.1f}%" if nd else "",
              "warp", round(wn), round(c_warp[cno]), f"{(c_warp[cno] - wn) / wn * 100:+.1f}%" if wn else "")
    tw = sum(c_prod[c] * warp_m[c] for c in cons); tws = sum(c_warp.values())
    print(f"WARP total: needed {tw:,.0f} kg vs sent to sizing {tws:,.0f} kg ({(tws - tw) / tw * 100:+.1f}%)")
    print("other contracts:", other)
    print("flagged days:", len(flags))
    for r in sorted(flags, key=lambda r: (-r[4], -r[1]))[:40]:
        print(f"{r[0]} | {r[1]:,.0f} kg | woven {r[2]:,} m | {r[3]}")
    print("SIZING:")
    for r in sorted(sizing, key=lambda r: -r[5]):
        print(r)
    bad_weeks = [r for r in daily if r[6] > 400 and r[4] > 1.3 * max(r[5], 1)]
    print("weeks with excess:", len(bad_weeks), [(str(r[0]), r[4], r[5]) for r in bad_weeks[:30]])
    print("saved", out)


if __name__ == "__main__":
    main()
