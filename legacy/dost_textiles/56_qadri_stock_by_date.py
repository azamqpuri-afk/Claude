"""
56_qadri_stock_by_date.py - Qadri Traders' 52/1 & 54/1 cotton (codes 6/18, contracts 38,50,60,68,70): the stock the mill
holds for Qadri, in bags, on month-end dates and today - fresh fabric + warp on beams on looms + sized beams waiting +
yarn at the sizer + yarn bags in hand. ERP read only. Separate Google Sheet "Qadri 52-54 ctn stock in bags by date".
Upload:  python 56_qadri_stock_by_date.py --upload <google sheet id>
"""
import importlib.util
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
_s = importlib.util.spec_from_file_location("s49", HERE / "49_contract_closing_by_sets.py")
s49 = importlib.util.module_from_spec(_s); _s.loader.exec_module(s49)
cc = s49.cc; R = cc.dash.rows
OUT = HERE / "financial_statements" / "qadri_52_54_stock_by_date.xlsx"
CL = "38,50,60,68,70"
TODAY = datetime.combine(datetime.today().date(), datetime.min.time())
DATES = [datetime(2026, 6, 30), datetime(2026, 7, 31), datetime(2026, 8, 31), TODAY]
PHYSICAL_TODAY = 68          # owner's count, 28-Sep-2026
REJ_TODAY = 45000            # owner: rejection in stock today (mill's)
NAVY = PatternFill("solid", fgColor="1F3B57"); INPUT = PatternFill("solid", fgColor="FFF4C2"); GREEN = PatternFill("solid", fgColor="D9F2E3")
W = Font(bold=True, color="FFFFFF"); B = Font(bold=True); T = Font(bold=True, size=15, color="1F3B57"); S = Font(bold=True, size=12, color="1F3B57")
NOTE = Font(italic=True, color="4B4F5A"); N0, N1, N4 = '#,##0;[Red]-#,##0', '#,##0.0;[Red]-#,##0.0', '0.0000'
BAGEQ = ("CASE WHEN NVL(d.bags,0) = 0 AND NVL(d.cones,0) = 0 THEN NVL(d.yarn_wt,0)/45.36 "
         "ELSE NVL(d.bags,0) + NVL(d.cones,0)/NVL(NULLIF(d.cone_pbag,0), CASE WHEN y.count > 40 THEN 40 ELSE 24 END) END")   # weight-only lines -> kg / bag


def head(ws, r, labels):
    for i, h in enumerate(labels, 1):
        c = ws.cell(row=r, column=i, value=h); c.font, c.fill = W, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


# one set of tabs per yarn. VA: 36/2 is two-ply, weighed as count 18 (owner confirmed 30-Sep-2026; the sized beams carry
# 756 kg per 9,600 set m = 0.1736 lb/m against 0.1747 at count 18). No physical count / rejection figure given for VA -> ERP.
YARNS = [
    dict(key="ctn", CL=CL, YC="6,18", SHRINK_CL=CL + ",52", TAB="Stock by date", BTAB="Beams by date",
         TITLE="QADRI TRADERS · 52/1 & 54/1 COTTON · STOCK HELD BY THE MILL, IN BAGS",
         WARP="=90.43*46*1.0936/800/51.5", WEFT="=82*46*1.0936/800/52.5", PHYS=PHYSICAL_TODAY, REJ=REJ_TODAY),
    dict(key="va", CL="67", YC="98", SHRINK_CL="67", TAB="36-2 VA stock by date", BTAB="36-2 VA beams by date",
         TITLE="QADRI TRADERS · 36/2 VA (contract 67) · STOCK HELD BY THE MILL, IN BAGS",
         WARP="=50*46*1.0936/800/18", WEFT="=42*46*1.0936/800/18", PHYS=None, REJ=None),
]


def main():
    cur = cc.dash.connect(cc.dash.load_config()).cursor()
    wb = Workbook()
    for i, P in enumerate(YARNS):
        build(cur, wb, P, i == 0)
    method(wb)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)


def build(cur, wb, P, first):
    CL, YC = P["CL"], P["YC"]
    cids = {cid for (cid,) in R(cur, f"SELECT cont_id FROM dostex.contm WHERE cont_no IN ({CL})")}
    prod = R(cur, f"""SELECT TRUNC(m.td), SUM(NVL(d.a,0)+NVL(d.a2,0)), SUM(NVL(d.b,0)), SUM(NVL(d.c,0)+NVL(d.cp,0))
        FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no=m.trans_no
        WHERE m.tt='PRODUCTION' AND d.cont_id IN (SELECT cont_id FROM dostex.contm WHERE cont_no IN ({CL})) GROUP BY TRUNC(m.td)""")
    gp = R(cur, f"""SELECT TRUNC(v.trans_date), v.trans_type, SUM(NVL(v.a,0)+NVL(v.a2,0)), SUM(NVL(v.b,0)), SUM(NVL(v.c,0)+NVL(v.cp,0)),
        CASE WHEN v.trans_type = 'GPOUT' AND UPPER(NVL(v.deliver_to,' ')) LIKE '%REJECTION%' THEN 1 ELSE 0 END, SUM(v.total_mtr)
        FROM dostex.fab_gp_v v JOIN dostex.wc_cont_sub_v s ON s.cont_sub_id=v.cont_sub_id WHERE s.cont_no IN ({CL})
        GROUP BY TRUNC(v.trans_date), v.trans_type, CASE WHEN v.trans_type = 'GPOUT' AND UPPER(NVL(v.deliver_to,' ')) LIKE '%REJECTION%' THEN 1 ELSE 0 END""")
    mount = {(t, s_): (float(w or 3), td, (ln or "").strip()) for t, s_, w, td, ln in R(cur, s49.MOUNT_SQL)}
    # every mounting of a beam (a beam can move from one loom to another, e.g. 00550 L13 -> L33 on 10-Jul-2026):
    # the loom shown for a date is the last one it was mounted on by that date
    mounts = defaultdict(list)
    for t, s_, td, ln in R(cur, """SELECT d.ref_tno, d.ref_sr, TRUNC(m.td), l.lname FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no = m.trans_no
                                   LEFT JOIN dostex.looms l ON l.lcode = d.lcode WHERE m.tt = 'BEAMS_TO_LOOM'"""):
        mounts[(t, s_)].append((td, (ln or "").strip()))
    off = {(t, s_): td for t, s_, st, td in R(cur, s49.OFF_SQL)}
    offst = {(t, s_): st for t, s_, st, td in R(cur, s49.OFF_SQL)}
    bprod = defaultdict(list)
    for t, s_, m, d in R(cur, """SELECT d.ref_tno, d.ref_sr, NVL(d.a,0)+NVL(d.a2,0)+NVL(d.b,0)+NVL(d.c,0)+NVL(d.cp,0), TRUNC(m.td)
            FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no=m.trans_no WHERE m.tt='PRODUCTION'"""):
        bprod[(t, s_)].append((d, float(m or 0)))
    cno = {cid: int(n) for cid, n in R(cur, f"SELECT cont_id, cont_no FROM dostex.contm WHERE cont_no IN ({CL})")}
    beams = [(tr, sr, setno, bcode, float(blen or 0), cno[cid], td) for tr, sr, setno, bcode, blen, cid, tt, tno, g, td, *_ in R(cur, s49.BEAMS_SQL)
             if cid in cids]
    mv = R(cur, f"""SELECT TRUNC(m.td), m.tt, c.cont_no, a.an, SUM({BAGEQ}), SUM({cc.KG}) FROM dostex.wytm m JOIN dostex.wytd d ON d.trans_no=m.trans_no
        LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id LEFT JOIN dostex.ach a ON a.ac=NVL(m.ac1,m.ac2)
        WHERE d.ycode IN ({YC}) AND c.cont_no IN ({CL}) AND NOT (m.tt='OGP_YARN' AND REGEXP_LIKE(NVL(m.tdesc,' '),'@ *[0-9]'))
        GROUP BY TRUNC(m.td), m.tt, c.cont_no, a.an""")
    SIGN = {"IGP_YARN": 1, "OGP_YARN": -1, "TO_SIZING": -1, "FROM_SIZING": 1, "YARN_TO_LOOMS": -1, "YARN_FROM_LOOMS": 1}

    # same basis as the grand closing (55_): shrinkage on finished beams and sized warp per set meter (contracts incl. Azam Puri 52,
    # whose set was sized from contract 50's yarn); fabric still to come = total set length x (1 - shrinkage) - fabric woven
    cids52 = {cid for (cid,) in R(cur, f"SELECT cont_id FROM dostex.contm WHERE cont_no IN ({P['SHRINK_CL']})")}
    b52 = [(tr, sr, float(blen or 0)) for tr, sr, setno, bcode, blen, cid, *_ in R(cur, s49.BEAMS_SQL) if cid in cids52]
    done_len = sum(bl * mount.get((tr, sr), (3.0, None, ""))[0] for tr, sr, bl in b52 if offst.get((tr, sr)) == "COMPLETE")
    done_wov = sum(sum(m for d, m in bprod[(tr, sr)]) for tr, sr, bl in b52 if offst.get((tr, sr)) == "COMPLETE")
    SHRINK = 1 - done_wov / done_len
    all_len = sum(bl * mount.get((tr, sr), (3.0, None, ""))[0] for tr, sr, bl in b52)
    BQ2 = "NVL(d.bags,0) + NVL(d.cones,0)/NVL(NULLIF(d.cone_pbag,0), CASE WHEN y.count > 40 THEN 40 ELSE 24 END)"
    sent_b = sum(float(v or 0) * (1 if tt == "TO_SIZING" else -1) for tt, v in R(cur, f"""SELECT m.tt, SUM({BQ2}) FROM dostex.wytm m
        JOIN dostex.wytd d ON d.trans_no=m.trans_no LEFT JOIN dostex.yarn y ON y.ycode=d.ycode JOIN dostex.contm c ON c.cont_id=d.cont_id
        WHERE m.tt IN ('TO_SIZING','FROM_SIZING') AND d.ycode IN ({YC}) AND c.cont_no IN ({P['SHRINK_CL']}) GROUP BY m.tt"""))
    # loom production by day (all contracts) for each loom's daily rate, and per loom+contract for the last production check
    lprod = defaultdict(list); lcprod = defaultdict(list)
    for ln_, cn_, d_, m_ in R(cur, """SELECT l.lname, p.cont_no, TRUNC(p.td), SUM(p.total_mtr) FROM dostex.wv_prod_v p JOIN dostex.looms l ON l.lcode=p.lcode
                                     GROUP BY l.lname, p.cont_no, TRUNC(p.td)"""):
        lprod[(ln_ or "").strip()].append((d_, float(m_ or 0))); lcprod[((ln_ or "").strip(), int(cn_))].append(d_)

    def unentered(D, bs):
        # fabric woven on the looms but not yet entered on date D (same estimate as the grand closing): per loom + contract, days since
        # the last production entry on or after mounting (or since mounting) x the loom's average meters a day over the 30 days to D
        grp, cap = {}, defaultdict(float)
        for r_ in bs:
            if r_[8]:
                k = (r_[3], r_[0]); grp[k] = min(grp.get(k, r_[9]), r_[9])
                cap[k] += max(0.0, r_[10] * (1 - SHRINK) - r_[11])          # the most fabric its beams can still make
        tot = 0.0
        for (ln_, cn_), md_ in grp.items():
            last = max([d_ for d_ in lcprod.get((ln_, cn_), []) if md_ <= d_ <= D], default=md_)
            # the loom wove ANOTHER contract after this one's last entry: this beam is not running (take-off keyed late)
            if any(last < d_ <= D for (l2, c2), ds in lcprod.items() if l2 == ln_ and c2 != cn_ for d_ in ds):
                continue
            rate = sum(m_ for d_, m_ in lprod.get(ln_, []) if D - timedelta(days=30) < d_ <= D) / 30
            tot += min((D - last).days * rate, cap[(ln_, cn_)])
        return tot

    def fab(D):
        a = b = c = 0.0
        for d, pa, pb, pc in prod:
            if d <= D:
                a += float(pa or 0); b += float(pb or 0); c += float(pc or 0)
        for d, tt, ga, gb, gc, _, _t in gp:
            if d <= D:
                sg = -1 if tt == "GPOUT" else 1
                a += sg * float(ga or 0); b += sg * float(gb or 0); c += sg * float(gc or 0)
        return a, b, c

    def dispatched(D):
        # fabric dispatched to Qadri: the mill's own rejection sale (Khalid) is not Qadri's
        return sum((1 if tt == "GPOUT" else -1) * float(t or 0) for d, tt, ga, gb, gc, rej, t in gp if d <= D and not rej)

    def received(D):
        return sum(SIGN[tt] * float(bq or 0) for d, tt, cn, an, bq, kg in mv if d <= D and tt in ("IGP_YARN", "OGP_YARN"))

    def book_yarn(D):
        return sum(SIGN.get(tt, 0) * float(bq or 0) for d, tt, cn, an, bq, kg in mv if d <= D)

    def at_sizer(D):
        bal = defaultdict(float); last = {}
        for d, tt, cn, an, bq, kg in mv:
            if d > D or tt not in ("TO_SIZING", "FROM_SIZING", "IGP_SIZED_BEAMS"):
                continue
            k = (cn, (an or "").strip())
            bal[k] += float(bq or 0) * 45.36 * (1 if tt == "TO_SIZING" else -1)     # bags + cones as part-bags (owner counts bags)
            if tt == "TO_SIZING":
                last[k] = max(last.get(k, d), d)
        return [(k, v) for k, v in bal.items() if v >= 3 * 45.36 and last.get(k) and last[k] >= D - timedelta(days=60)]

    def beam_state(D):
        out = []
        for tr, sr, setno, bcode, blen, cn, td in beams:
            if td > D:
                continue
            mw, md, ln = mount.get((tr, sr), (3.0, None, ""))
            ln = max((x for x in mounts.get((tr, sr), []) if x[0] <= D), default=(None, ln))[1]   # the loom it was on at D
            od = off.get((tr, sr))
            if od and od <= D:
                continue
            alen = blen * mw
            woven = sum(m for d, m in bprod[(tr, sr)] if d <= D)
            on = bool(md and md <= D)
            out.append([cn, setno, bcode, ln if on else "Waiting (not mounted)", f"{td:%d-%b-%y}", round(alen), round(woven if on else 0),
                        round(max(0.0, alen - woven) if on else alen), on, md, alen, woven if on else 0.0])
        return sorted(out, key=lambda r: (not r[8], r[0], r[4], str(r[2])))

    KGM = (sent_b * 45.36 - sum(v for _, v in at_sizer(TODAY))) / all_len     # sized warp kg per set meter, as in the grand closing

    rej_now = fab(TODAY)[2]
    # 60/1 cotton weft is a separate account (owner, 30-Sep-2026): fabric with 60/1 weft, all dispatched before the first date
    m60 = float(R(cur, """SELECT SUM(CASE WHEN v.trans_type = 'GPOUT' THEN v.total_mtr ELSE -v.total_mtr END) FROM dostex.fab_gp_v v JOIN dostex.wc_cont_sub_v s ON s.cont_sub_id=v.cont_sub_id
                          WHERE s.cont_no IN (38,50,60,68,70) AND s.quality LIKE '%54/60%'
                              AND NOT (v.trans_type = 'GPOUT' AND UPPER(NVL(v.deliver_to,' ')) LIKE '%REJECTION%')""")[0][0] or 0) if P["key"] == "ctn" else 0.0
    TAB = P["TAB"]
    if first:
        ws = wb.active; ws.title = TAB
    else:
        ws = wb.create_sheet(TAB, 1)
    ws["A1"] = P["TITLE"]; ws["A1"].font = T
    ws["A2"] = (f"Contracts {CL}. Fresh fabric + warp on beams + sized beams waiting + yarn at sizer + yarn bags in hand. "
                f"1 bag = 100 lb. Made {TODAY:%d-%b-%Y}; ERP read only."); ws["A2"].font = NOTE
    ws.column_dimensions["A"].width = 58
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["G"].width = 60
    hand_now = book_yarn(TODAY)
    ins = [("Warp weight (lb / m)", P["WARP"]), ("Weft weight (lb / m)", P["WEFT"]),
           ("Fabric weight (lb / m)", "=B4+B5"), ("Lb per bag", 100), ("Kg per bag", 45.36),
           ("Rate (Rs per bag)", 61000 if P["key"] == "ctn" else None),
           ("Rejection in stock today - mill's (m)" + ("" if P["REJ"] else " - ERP, no count given"), P["REJ"] if P["REJ"] else round(rej_now)),
           # the count is rolled forward with ERP receipts / issues after the count date (28-Sep-2026)
           ("Yarn bags in hand today - physical count 28-Sep + ERP movements since" if P["PHYS"] else "Yarn bags in hand today - ERP (no physical count given)",
            f"={P['PHYS']}+{round(hand_now - book_yarn(datetime(2026, 9, 28)), 2)}" if P["PHYS"] else round(hand_now, 1))]
    for i, (l, v) in enumerate(ins, 4):
        ws.cell(row=i, column=1, value=l); c = ws.cell(row=i, column=2, value=v)
        c.number_format = N4 if i <= 6 else N0
        if i >= 7:
            c.fill = INPUT
    # B4 warp B5 weft B6 fabric B7 lb/bag B8 kg/bag B9 rate B10 rej today B11 physical today · E4 shrinkage · E5 kg per set meter
    ws["D4"] = "Shrinkage beam → fabric (finished beams)"; ws["E4"] = round(SHRINK, 4); ws["E4"].number_format = "0.00%"; ws["E4"].fill = INPUT
    ws["D5"] = "Sized warp kg per meter of set"; ws["E5"] = round(KGM, 5); ws["E5"].number_format = "0.00000"
    ws.column_dimensions["D"].width = 34
    hr = 13
    head(ws, hr, ["As on"] + [("Today " if D == TODAY else "") + f"{D:%d-%b-%y}" for D in DATES] + ["", "How"])
    L = "BCDE"
    rows_ = []
    for D in DATES:
        a, b, c = fab(D)
        bs = beam_state(D)
        rows_.append(dict(total_fab=a + b + c, erp_rej=c, loom_m=sum(r[7] for r in bs if r[8]), nloom=sum(1 for r in bs if r[8]),
                          loom_set=sum(r[10] for r in bs if r[8]), loom_wov=sum(r[11] for r in bs if r[8]), unent=unentered(D, bs),
                          wait_m=sum(r[7] for r in bs if not r[8]), nwait=sum(1 for r in bs if not r[8]),
                          sizer_b=sum(v for _, v in at_sizer(D)) / 45.36, erp_yarn=book_yarn(D), recv=received(D), disp=dispatched(D), beams=bs))
    lines = [
        ("FABRIC", None, None),
        ("Fabric in stock, all grades (ERP, m)", "total_fab", "production less gate passes up to the date"),
        ("Less: rejection in stock - mill's (m)", "rej", "today: your figure; earlier dates rolled back with ERP C/CP movements"),
        ("Fresh fabric for Qadri (m)", "=r15-r16", ""),
        ("→ Fresh fabric (bags)", "=r17*$B$6/$B$7", "× fabric weight"),
        ("WARP", None, None),
        ("Beams on looms (count)", "nloom", ""),
        ("Set length of beams on looms (m) · fabric already woven from them (m)", "loom_pair", "beam length × 3 pieces · production entered against those beams"),
        ("→ Warp on looms, as fabric still to come (bags)", "=(LOOMSET*(1-$E$4)-LOOMWOV)*$B$4/$B$7", "total set length × (1 - shrinkage) - fabric woven, × warp weight"),
        ("Fabric woven on looms, not yet entered (m) - estimate", "unent", "days since last production check × loom's meters a day (30 days to the date)"),
        ("→ Weft in it (bags)", "=r23*$B$5/$B$7", "× weft weight (its warp is in the row above)"),
        ("Sized beams waiting (count)", "nwait", ""),
        ("Sized beams waiting (m)", "wait_m", "beam length × 3 pieces"),
        ("→ Sized beams waiting, as fabric to come (bags)", "=r26*(1-$E$4)*$B$4/$B$7", "less shrinkage, × warp weight"),
        ("Yarn at sizer - physical (bags)", "sizer_b", "bags + cones sent less bags + cones back (as sized beams or loose); lots of 3+ bags sent within 60 days"),
        ("→ Yarn at sizer, as fabric to come (bags)", "=r28*$B$8/$E$5*(1-$E$4)*$B$4/$B$7", "bags × 45.36 kg ÷ sized warp per set meter, less shrinkage, × warp weight (closing weight)"),
        ("YARN", None, None),
        ("Yarn bags in hand (ERP stock on the date)", "erp_yarn", "received - returned - to sizing + back - to looms + back"),
        ("→ Yarn bags in hand used", "hand", "earlier dates: ERP stock · today: your physical count"),
        ("TOTAL STOCK HELD FOR QADRI (bags)", "=r18+r22+r24+r27+r29+r32", ""),
        ("Value (Rs)", "=r33*$B$9", ""),
        ("CHECK", None, None),
        ("Yarn received from Qadri, net (bags)", "recv", ""),
        ("Fabric dispatched to Qadri, net (m)", "disp", ""),
        ("Yarn in fabric dispatched (bags)", "=(r37*$B$6-M60*$B$5)/$B$7" if P["key"] == "ctn" else "=r37*$B$6/$B$7",
         f"{m60:,.0f} m of it (contract 38 OGP 361-381, Apr-Jun 2025) had 60/1 weft - a separate account, so only its warp counts here" if m60 else ""),
        ("Book balance = received - dispatched (bags)", "=r36-r38", "what Qadri's account says the mill should hold"),
        ("Difference: stock held - book (bags)", "=r33-r39", "+ = Qadri short-supplied (owes the mill) · − = mill holds that much more of Qadri's yarn (weft in fabric on looms not yet entered is part of it)"),
    ]
    for i, (lab, key, how) in enumerate(lines, hr + 1):
        ws.cell(row=i, column=1, value=lab)
        if key is None:
            ws.cell(row=i, column=1).font = S
            continue
        for j, D in enumerate(DATES):
            col = L[j]; rw = rows_[j]
            if key.startswith("="):
                v = "".join(col + p if k else p for k, p in _split(key)).replace("LOOMSET", str(round(rw["loom_set"]))).replace("LOOMWOV", str(round(rw["loom_wov"]))).replace("M60", str(round(m60)))
            elif key == "loom_pair":
                v = f"{rw['loom_set']:,.0f} · {rw['loom_wov']:,.0f}"
            elif key == "rej":
                v = f"={round(rw['erp_rej'])}+$B$10-{round(rej_now)}"
            elif key == "hand":
                v = f"=$B$11" if D == TODAY else f"={col}{i - 1}"
            else:
                v = round(rw[key], 1)
            c = ws.cell(row=i, column=j + 2, value=v)
            c.number_format = N0 if ("(m)" in lab or "count" in lab or "Rs" in lab or "kg" in lab) else N1
        ws.cell(row=i, column=7, value=how).font = NOTE
        if lab.startswith("TOTAL") or lab.startswith("→") or lab.startswith("Book") or lab.startswith("Difference"):
            for col in range(1, 6):
                ws.cell(row=i, column=col).font = B
        if lab.startswith("TOTAL"):
            for col in range(1, 6):
                ws.cell(row=i, column=col).fill = GREEN
    ws.freeze_panes = "B14"

    bd = wb.create_sheet(P["BTAB"])
    bd["A1"] = f"Beams on looms and sized beams waiting, on each date · contracts {CL}"; bd["A1"].font = T
    r = 3
    for D, rw in zip(DATES, rows_):
        bd.cell(row=r, column=1, value=f"As on {D:%d-%b-%Y}").font = S; r += 1
        head(bd, r, ["Contract", "Set", "Beam", "Loom", "Received", "Set length (m)", "Woven from beam (m)", "Unwoven (m)", "Fabric still to come (m)", "Warp (bags)"]); r += 1
        f0 = r
        for b_ in rw["beams"]:
            for i, v in enumerate(b_[:8], 1):
                bd.cell(row=r, column=i, value=v)
            # fabric still to come = total set length x (1 - shrinkage) - fabric woven; warp bags on that
            bd.cell(row=r, column=9, value=f"=MAX(0,F{r}*(1-'{TAB}'!$E$4)-G{r})").number_format = N0
            bd.cell(row=r, column=10, value=f"=I{r}*'{TAB}'!$B$4/'{TAB}'!$B$7").number_format = N1
            r += 1
        bd.cell(row=r, column=1, value="TOTAL").font = B
        for col in "FGHIJ":
            c = bd[f"{col}{r}"]; c.value = f"=SUM({col}{f0}:{col}{r - 1})"; c.font = B; c.number_format = N1 if col == "J" else N0
        for (cn_, an), kg in at_sizer(D):
            r += 1; bd.cell(row=r, column=1, value=f"At sizer: {an}, contract {cn_}: {kg / 45.36:,.1f} bags physical ({kg:,.0f} kg)")
        r += 3
    for i, w in enumerate([11, 10, 9, 22, 11, 13, 14, 12, 14, 11], 1):
        bd.column_dimensions[bd.cell(row=1, column=i).column_letter].width = w
    for D, rw in zip(DATES, rows_):
        print(P["key"], D.date(), {k: round(v, 1) for k, v in rw.items() if k != "beams"})


def method(wb):
    m = wb.create_sheet("Method")
    for i, t in enumerate([
        "HOW THIS IS MADE",
        "Only Qadri Traders' 52/1 & 54/1 cotton (same yarn, ERP codes 6 and 18), contracts 38, 50, 60, 68, 70.",
        "Fabric in stock on a date = production entered up to the date less fabric gate passes (out less returns) up to the date, by grade.",
        "Rejection is the mill's: today's 45,000 m is your figure; for earlier dates it is rolled back with the ERP rejection (C/CP) movements.",
        "Warp on looms = each beam mounted by the date and not taken off: beam length × 3 pieces less fabric woven from that beam by the date.",
        "Sized beams waiting = beams received from the sizer by the date and not yet mounted.",
        "Yarn at sizer = yarn sent less taken back less sized beams received, per sizer and contract; only lots of 3+ bags sent in the last 60 days.",
        "Yarn bags in hand: ERP yarn stock on the earlier dates (bags + cones as part-bags); today your physical count.",
        "Conversion: fabric 0.2086 lb/m, warp 0.1104 lb/m (ends 90.43, picks 82, width 46, ÷800, counts 51.5 / 52.5); 1 bag = 100 lb.",
        "Same basis as the grand closing: warp on looms, waiting beams and the sizer lot are counted as the fabric they will make - shrinkage (E4, from finished beams) on the TOTAL set length, less fabric already woven; weft in fabric woven but not yet entered is estimated for each date (days since the last production check x the loom's meters a day over the 30 days to that date).",
        "Check rows: yarn received less yarn in the fabric dispatched gives the book balance; the difference is Qadri's short supply to date.",
        "",
        "36/2 VA TABS (contract 67, ERP yarn code 98) - same method, with these differences:",
        "Count: 36/2 is two-ply, weighed as count 18 for warp and weft (ends 50, picks 42, width 46). Check: the sized beams carry 756 kg per 9,600 m of set = 0.1736 lb/m vs 0.1747 at count 18.",
        "Shrinkage (E4) and sized warp per set meter (E5) are measured on contract 67's own finished beams and sizing.",
        "No physical yarn count or rejection figure given for VA: today's yarn in hand and rejection are the ERP figures (enter a count in B11 / B10 to replace them).",
        "Old-ERP leftovers on the earlier VA contracts (5, 15, 17, 49, about +3 bags of yarn and 175 m B-grade) are left out.",
    ], 1):
        m.cell(row=i, column=1, value=t).font = T if i == 1 else Font()
    m.column_dimensions["A"].width = 140


def _split(expr):
    """'=r14-r15' -> pieces; flag True where a column letter must be prefixed to the row number"""
    import re
    out = []
    for p in re.split(r"(r\d+)", expr):
        if re.fullmatch(r"r\d+", p):
            out.append((True, p[1:]))
        elif p:
            out.append((False, p))
    return out


def upload(fid):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    creds = service_account.Credentials.from_service_account_file(str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
    build("drive", "v3", credentials=creds).files().update(fileId=fid, media_body=MediaFileUpload(
        str(OUT), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
    print("uploaded", fid)


if __name__ == "__main__":
    main()
    if len(sys.argv) > 2 and sys.argv[1] == "--upload":
        upload(sys.argv[2])
