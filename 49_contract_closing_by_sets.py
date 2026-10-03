"""
49_contract_closing_by_sets.py
Contract closing by SET LENGTHS (2026-09-26) - a second Google Sheet in the
"Contracts closing" folder. The Oracle ERP is read, never changed.

Idea (owner): the sized beams (sets) are what produced the fabric, so the weft
that SHOULD have been used is the set length woven x weft weight per meter at the
owner's closing weft count. Compare that with the weft actually issued to looms
and with the weft yarn available, to see what the weft balance should be.

Per set (DOSTEX.WYBEAMS.SETNO, from IGP_SIZED_BEAMS / OPN_SIZED_BEAMS):
  beam length   = WYBEAMS.BLENGTH x mounting width multiplier (the ERP's own
                  GET_BEAM_ALENGTH logic: BEAMS_TO_LOOM row BWIDTH)
  fabric woven  = production rows pointing at the beam (WPRDTD REF_TNO/REF_SR =
                  WYBEAMS TRANS_NO/SR_NO; the ERP's own GET_BEAM_PROD logic)
  length basis for weft = full beam length for beams taken off COMPLETE;
                  fabric woven for beams CUT early, still on loom or not mounted
Per contract (weft account, kg):
  weft yarn available = yarn received - yarn returned - yarn sent to sizing (net)
  weft needed on sets = length basis x weft lb/m (owner's weft count)
  weft balance should be = available - needed
  weft issued to looms (net of weft returned from looms) -> actual balance
  over / (under) issue = issued - needed, valued at the owner's weft rate
Counts and rates are the same owner inputs as the first closing file (Inputs tab,
editable; weft weight per meter recalculates).
"""

import importlib.util
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
_s = importlib.util.spec_from_file_location("cc", HERE / "48_contract_closing.py")
cc = importlib.util.module_from_spec(_s); _s.loader.exec_module(cc)
dash = cc.dash
KG_PER_LB = cc.KG_PER_LB
OUT = HERE / "financial_statements" / "dost_contract_closing_by_sets.xlsx"

BEAMS_SQL = """
 SELECT b.trans_no, b.sr_no, b.setno, b.bcode, b.blength, d.cont_id, m.tt, m.tno, m.gp_no, m.td, m.tdesc,
        d.ends, d.set_length, a.an
 FROM dostex.wybeams b JOIN dostex.wytm m ON m.trans_no = b.trans_no
 JOIN dostex.wytd d ON d.trans_no = b.trans_no AND d.sr_no = b.ref_sr
 LEFT JOIN dostex.ach a ON a.ac = NVL(m.ac1, m.ac2)
 WHERE m.tt IN ('IGP_SIZED_BEAMS','OPN_SIZED_BEAMS')"""
MOUNT_SQL = """
 SELECT d.ref_tno, d.ref_sr, MAX(NVL(d.bwidth,1)), MIN(m.td), MAX(l.lname)
 FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no = m.trans_no LEFT JOIN dostex.looms l ON l.lcode = d.lcode
 WHERE m.tt = 'BEAMS_TO_LOOM' GROUP BY d.ref_tno, d.ref_sr"""
OFF_SQL = """
 SELECT d.ref_tno, d.ref_sr, MAX(d.bstatus), MAX(m.td)
 FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no = m.trans_no
 WHERE m.tt = 'BEAMS_FROM_LOOM' GROUP BY d.ref_tno, d.ref_sr"""
PROD_SQL = """
 SELECT d.ref_tno, d.ref_sr, d.cont_id, SUM(NVL(d.a,0)+NVL(d.a2,0)+NVL(d.b,0)+NVL(d.c,0)+NVL(d.cp,0)), MIN(m.td), MAX(m.td)
 FROM dostex.wprdtm m JOIN dostex.wprdtd d ON d.trans_no = m.trans_no
 WHERE m.tt = 'PRODUCTION' GROUP BY d.ref_tno, d.ref_sr, d.cont_id"""

NAVY, INPUT, SEC, GREEN, AMBER = cc.NAVY, cc.INPUT, cc.SEC, cc.GREEN, cc.AMBER
WHITE, B = cc.WHITE, cc.B
NUM0, NUM1, NUM4, PCT = cc.NUM0, cc.NUM1, cc.NUM4, cc.PCT
RED = PatternFill("solid", fgColor="F8C9C9")


def hdr(ws, row, labels, col0=1):
    for i, h in enumerate(labels, col0):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = WHITE, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32


def main():
    cur = dash.connect(dash.load_config()).cursor()
    rows = dash.rows
    contracts = sorted(rows(cur, cc.CONTRACTS_SQL), key=lambda c: int(c[1]))
    cinfo = {c[0]: c for c in contracts}
    yarn = {}
    for cid, tt, kg, nowt, bg, cn in rows(cur, cc.YARN_SQL):
        yarn[(cid, tt)] = float(kg or 0)
    mount = {(t, s): (float(w or 1), td, ln) for t, s, w, td, ln in rows(cur, MOUNT_SQL)}
    off = {(t, s): (st, td) for t, s, st, td in rows(cur, OFF_SQL)}
    prod = {}
    for t, s, cid, m, f, l in rows(cur, PROD_SQL):
        prod.setdefault((t, s), []).append((cid, float(m or 0), f, l))
    # sized actual kg per receipt (description weight), as in the first file
    sized_rows = rows(cur, cc.SIZED_SQL)
    desc_by_trans = {}
    for tr, cid, tno, gp, td, desc, ywt, skg in sized_rows:
        desc_by_trans[tr] = cc.desc_kg(desc)

    # build sets: key (contract, setno)
    sets = {}
    for tr, sr, setno, bcode, blen, cid, tt, tno, gp, td, desc, ends, slen, sizer in rows(cur, BEAMS_SQL):
        mw, mdate, loom = mount.get((tr, sr), (1.0, None, None))
        stt, odate = off.get((tr, sr), (None, None))
        plist = prod.get((tr, sr), [])
        woven = sum(p[1] for p in plist)
        other = sum(p[1] for p in plist if p[0] != cid)
        alen = float(blen or 0) * mw
        if stt == "COMPLETE":
            status, basis = "Complete", alen
        elif stt == "CUT":
            status, basis = "Cut early", woven
        elif mdate:
            status, basis = "On loom", woven
        else:
            status, basis = "Not mounted", woven
        s = sets.setdefault((cid, setno), dict(cid=cid, setno=setno, tt=tt, tno=tno, gp=gp, td=td, desc=desc, ends=ends,
                                                slen=slen, sizer=sizer, beams=0, comp=0, cut=0, run=0, idle=0,
                                                alen=0.0, woven=0.0, other=0.0, basis=0.0, desc_kg=desc_by_trans.get(tr)))
        s["beams"] += 1
        s["alen"] += alen; s["woven"] += woven; s["other"] += other; s["basis"] += basis
        s[{"Complete": "comp", "Cut early": "cut", "On loom": "run", "Not mounted": "idle"}[status]] += 1

    wb = Workbook()
    # ---------------- Inputs (same owner counts/rates as file 1)
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = "Weft closing counts and rates (same as the first closing file; edit yellow cells - every contract recalculates)"
    inp["A1"].font = Font(bold=True, size=13, color="1F3B57")
    inp["A2"] = "Allowance % on weight:"
    inp["H2"] = 0
    inp["H2"].fill, inp["H2"].number_format = INPUT, "0.0%"
    cc.header(inp, 5, ["Key", "Yarn", "Party (blank = all)", "ERP yarn code", "Warp closing count", "Weft closing count",
                       "Rate Rs per lb", "Note"], [14, 24, 18, 10, 12, 12, 12, 40])
    entries = [(yc, "", w, f, rate) for yc, (_, w, f, rate) in cc.YARNS.items()]
    entries += [(yc, p, w, f, cc.YARNS[yc][3]) for (yc, p), (w, f) in cc.OVERRIDES.items()]
    used = {c[15] for c in contracts} - set(cc.YARNS)
    entries += [(yc, "", None, None, None) for yc in sorted(used) if yc]
    r = 6
    for yc, p, w, f, rate in entries:
        for col, v in enumerate([f"{yc}|{p}", cc.YARNS[yc][0] if yc in cc.YARNS else f"code {yc}", p, yc, w, f, rate,
                                 "" if yc in cc.YARNS else "NEEDED - weft count and rate not yet given"], 1):
            c = inp.cell(row=r, column=col, value=v)
            if col in (5, 6, 7):
                c.fill = INPUT
        r += 1
    last_inp = r - 1

    def key_for(yc, party):
        for (y, p) in cc.OVERRIDES:
            if y == yc and p in (party or ""):
                return f"{yc}|{p}"
        return f"{yc}|"

    # ---------------- Summary
    sm = wb.create_sheet("All contracts", 0)
    sm["A1"] = f"Contract closing by SET LENGTHS - weft that should have been used vs weft issued ({date.today():%d-%b-%Y})"
    sm["A1"].font = Font(bold=True, size=14, color="1F3B57")
    sm["A2"] = ("Weft needed = set length woven x weft lb/m at the owner's weft count. + over-issue = more weft issued to looms than "
                "the sets needed; - = less. ERP data read only; closing in this sheet.")
    sm["A2"].font = Font(italic=True, size=9)
    scols = ["Contract", "Party", "Status", "Weft yarn", "Sets", "Beams", "Set length (m)", "Fabric woven from sets (m)", "Beam shortage %",
             "Length basis for weft (m)", "Weft count", "Weft lb / m", "Weft needed on sets (kg)",
             "Yarn received net (kg)", "Yarn sent to sizing net (kg)", "Weft yarn available (kg)",
             "Weft balance SHOULD BE (kg)", "Weft issued to looms (kg)", "Actual weft balance (kg)",
             "Over / (under) issue vs sets (kg)", "Over / (under) issue %", "Over / (under) issue (Rs)", "Flags",
             "Weft needed on fabric woven (kg)", "Over / (under) vs fabric woven (kg)", "Over / (under) % vs fabric woven"]
    cc.header(sm, 4, scols, [8, 22, 16, 18, 6, 7, 12, 12, 9, 12, 8, 9, 12, 12, 12, 12, 13, 12, 12, 13, 9, 13, 60, 13, 13, 10])
    srow = 5
    tab_names = set()
    order = sorted({s["cid"] for s in sets.values() if s["cid"] in cinfo}, key=lambda c: int(cinfo[c][1]))
    for cid in order:
        (_, no, party, cdate, exp, cqty, f, l, mtr, ag, rj, q, wy, ends, cww, fy, picks, cfw, width, conv) = cinfo[cid]
        party = party or ""
        if int(no) in cc.WEFT_MIX:
            picks = sum(p for *_, p in cc.WEFT_MIX[int(no)])
        fy = cc.WEFT_CODE_FOR.get(int(no), fy)
        finished = l.date() <= cc.CUTOFF
        cs_sets = sorted([s for s in sets.values() if s["cid"] == cid], key=lambda s: (s["td"], str(s["setno"])))
        short = re.sub(r"[^A-Z0-9 ]", "", party.upper()).replace("TRADERSEMB", "").split(" ")[0] or "PARTY"
        name = f"{no} {short}"[:28]
        while name in tab_names:
            name += "."
        tab_names.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = "2E8B57" if finished else "E07B00"
        ws["A1"] = (f"Contract {no} - {party} - CLOSING BY SET LENGTHS" if finished else
                    f"Contract {no} - {party} - PENDING CLOSING (still running, last production {l:%d-%b-%Y})")
        ws["A1"].font = Font(bold=True, size=14, color="1F3B57" if finished else "9A3412")
        fk = key_for(fy, party)
        lk = lambda col: f"IFERROR(INDEX(Inputs!${col}$6:${col}${last_inp},MATCH(\"{fk}\",Inputs!$A$6:$A${last_inp},0)),\"\")"  # noqa: E731
        info = [("Quality", q), ("Weft yarn", cc.YARNS[fy][0] if fy in cc.YARNS else f"code {fy}"), ("Picks / inch", picks),
                ("Width (in)", width), ("Weft closing count", "=" + lk("F")), ("Weft rate / lb", "=" + lk("G")),
                ("Weft lb / m", "=IF(N(B8)>0,B6*B7*1.0936/(840*B8)*(1+Inputs!$H$2),\"\")"), ("Meters woven (all production)", mtr)]
        for i, (lab, v) in enumerate(info, 3):
            ws.cell(row=i + 1, column=1, value=lab)
            c = ws.cell(row=i + 1, column=2, value=v)
            c.number_format = {7: NUM4 if False else "General"}.get(i, "General")
        ws["B10"].number_format = NUM4
        ws["B11"].number_format = NUM0
        # set table
        hrow = 13
        ws.cell(row=hrow - 1, column=1, value="SETS (sized beams) AND THE FABRIC WOVEN FROM THEM").font = Font(bold=True, color="1F3B57")
        hdr(ws, hrow, ["Set #", "Received", "Gate pass #", "Sizer", "Ends", "Beams", "Complete", "Cut early", "On loom",
                       "Not mounted", "Set length (m)", "Fabric woven (m)", "Shortage %", "Woven but booked to another contract (m)",
                       "Length basis for weft (m)", "Weft needed (lb)", "Weft needed (kg)", "Sized weight (kg, description)"])
        r = hrow + 1
        first = r
        for s in cs_sets:
            vals = [s["setno"], s["td"], s["gp"], s["sizer"], s["ends"], s["beams"], s["comp"], s["cut"], s["run"], s["idle"],
                    round(s["alen"]), round(s["woven"]), None, round(s["other"]), round(s["basis"]), None, None, s["desc_kg"]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                if i == 2:
                    c.number_format = "dd-mmm-yy"
                elif i in (11, 12, 14, 15, 18):
                    c.number_format = NUM0
            ws.cell(row=r, column=13, value=f"=IF(K{r}>0,1-L{r}/K{r},\"\")").number_format = PCT
            ws.cell(row=r, column=16, value=f"=IFERROR(O{r}*$B$10,\"\")").number_format = NUM0
            ws.cell(row=r, column=17, value=f"=IFERROR(P{r}*{KG_PER_LB},\"\")").number_format = NUM0
            if s["other"] > 0:
                ws.cell(row=r, column=14).fill = AMBER
            r += 1
        last = r - 1
        ws.cell(row=r, column=1, value="TOTAL").font = B
        for col in (6, 7, 8, 9, 10, 11, 12, 14, 15, 16, 17, 18):
            L = get_column_letter(col)
            c = ws.cell(row=r, column=col, value=f"=SUM({L}{first}:{L}{last})" if last >= first else 0)
            c.font, c.number_format = B, NUM0
        tot = r
        ws.cell(row=tot, column=13, value=f"=IF(K{tot}>0,1-L{tot}/K{tot},\"\")").number_format = PCT
        # weft account
        r = tot + 2
        ws.cell(row=r, column=1, value="WEFT ACCOUNT (kg)").font = Font(bold=True, color="1F3B57")
        recv = yarn.get((cid, "IGP_YARN"), 0) - yarn.get((cid, "OGP_YARN"), 0)
        tosz = yarn.get((cid, "TO_SIZING"), 0) - yarn.get((cid, "FROM_SIZING"), 0)
        weft = yarn.get((cid, "YARN_TO_LOOMS"), 0) - yarn.get((cid, "YARN_FROM_LOOMS"), 0)
        acc = [("Yarn received from party, net of returns", round(recv)), ("Less: yarn sent to sizing (warp), net", round(tosz)),
               ("= Weft yarn available", f"=B{r + 1}-B{r + 2}"), ("Weft needed on set lengths", f"=Q{tot}"),
               ("= WEFT BALANCE SHOULD BE", f"=B{r + 3}-B{r + 4}"), ("Weft actually issued to looms (net)", round(weft)),
               ("= Actual weft balance", f"=B{r + 3}-B{r + 6}"), ("OVER / (UNDER) ISSUE vs set lengths", f"=B{r + 6}-B{r + 4}"),
               ("Over / (under) issue %", f"=IF(B{r + 4}>0,B{r + 8}/B{r + 4},\"\")"),
               ("Over / (under) issue value (Rs)", f"=IFERROR(B{r + 8}/{KG_PER_LB}*$B$9,\"\")"),
               ("Weft needed on FABRIC WOVEN from the sets", f"=IFERROR(L{tot}*$B$10*{KG_PER_LB},\"\")"),
               ("Over / (under) issue vs fabric woven", f"=B{r + 6}-B{r + 11}"),
               ("Over / (under) % vs fabric woven", f"=IF(N(B{r + 11})>0,B{r + 12}/B{r + 11},\"\")")]
        for i, (lab, v) in enumerate(acc, 1):
            ws.cell(row=r + i, column=1, value=lab)
            c = ws.cell(row=r + i, column=2, value=v)
            c.number_format = PCT if "%" in lab else NUM0
            if lab.isupper() or lab.startswith("= WEFT") or lab.startswith("OVER"):
                ws.cell(row=r + i, column=1).font = B
                c.font = B
        acc_row = r
        flags = []
        if cdate and cdate.date() <= date(2023, 10, 1):
            flags.append("Started before the new ERP: old-ERP yarn has no contract number - weft account incomplete")
        if any(s["tt"] == "OPN_SIZED_BEAMS" for s in cs_sets):
            flags.append("Includes sets already in the mill at go-live (opening sized beams)")
        if fy not in cc.YARNS:
            flags.append("Weft count/rate not given - fill Inputs")
        other = sum(s["other"] for s in cs_sets)
        if other > 100:
            flags.append(f"{other:,.0f} m woven from this contract's sets was booked to another contract")
        idle = sum(s["idle"] for s in cs_sets)
        if idle and finished:
            flags.append(f"{idle} beam(s) never mounted")
        ws.cell(row=2, column=1, value="; ".join(flags)).font = Font(italic=True, size=9, color="9A3412")
        for col, w in zip("ABCDEFGHIJKLMNOPQR", (40, 16, 10, 24, 8, 7, 9, 8, 8, 10, 12, 12, 9, 14, 13, 12, 12, 13)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = f"A{hrow + 1}"
        # summary row (links to the tab)
        T = f"'{name}'!"
        vals = [int(no), party, "Finished" if finished else "PENDING CLOSING", cc.YARNS[fy][0] if fy in cc.YARNS else f"code {fy}", len(cs_sets), f"={T}F{tot}", f"={T}K{tot}",
                f"={T}L{tot}", f"={T}M{tot}", f"={T}O{tot}", f"={T}B8", f"={T}B10", f"={T}Q{tot}",
                f"={T}B{acc_row + 1}", f"={T}B{acc_row + 2}", f"={T}B{acc_row + 3}", f"={T}B{acc_row + 5}", f"={T}B{acc_row + 6}",
                f"={T}B{acc_row + 7}", f"={T}B{acc_row + 8}", f"={T}B{acc_row + 9}", f"={T}B{acc_row + 10}", "; ".join(flags),
                f"={T}B{acc_row + 11}", f"={T}B{acc_row + 12}", f"={T}B{acc_row + 13}"]
        for i, v in enumerate(vals, 1):
            c = sm.cell(row=srow, column=i, value=v)
            c.number_format = {9: PCT, 12: NUM4, 21: PCT, 26: PCT, 23: "General"}.get(i, NUM0 if i >= 6 else "General")
        sm.cell(row=srow, column=3).fill = GREEN if finished else AMBER
        srow += 1
    sm.freeze_panes = "C5"
    sm.auto_filter.ref = f"A4:W{srow - 1}"
    # ---------------- By party and weft yarn (cross-contract bookings net out here)
    pv = wb.create_sheet("By party & weft yarn", 1)
    pv["A1"] = "Weft account by party and weft yarn - all contracts together (weft is often booked to the wrong contract of the same party)"
    pv["A1"].font = Font(bold=True, size=13, color="1F3B57")
    pv["A2"] = "Includes running contracts, because weft issued for a running contract is often booked to a finished one and vice versa."
    pv["A2"].font = Font(italic=True, size=9)
    cc.header(pv, 4, ["Party", "Weft yarn", "Contracts", "Set length (m)", "Weft needed on sets (kg)", "Weft yarn available (kg)",
                      "Weft balance SHOULD BE (kg)", "Weft issued (kg)", "Actual weft balance (kg)", "Over / (under) issue (kg)",
                      "Over / (under) %", "Over / (under) issue (Rs)", "Weft needed on fabric woven (kg)",
                      "Over / (under) vs fabric woven (kg)", "Over / (under) % vs fabric woven"], [24, 18, 9, 13, 14, 14, 14, 13, 13, 14, 10, 15, 14, 14, 11])
    A = "'All contracts'!"
    rng = lambda L: f"{A}${L}$5:${L}${srow - 1}"
    combos = []
    for row in range(5, srow):
        k = (sm.cell(row=row, column=2).value, sm.cell(row=row, column=4).value)
        if k not in combos:
            combos.append(k)
    for i, (p, y) in enumerate(sorted(combos, key=lambda k: (k[0] or "", k[1] or "")), 5):
        pv.cell(row=i, column=1, value=p); pv.cell(row=i, column=2, value=y)
        pv.cell(row=i, column=3, value=f"=COUNTIFS({rng('B')},A{i},{rng('D')},B{i})")
        for col, L in ((4, "G"), (5, "M"), (6, "P"), (7, "Q"), (8, "R"), (9, "S"), (10, "T"), (12, "V")):
            c = pv.cell(row=i, column=col, value=f"=SUMIFS({rng(L)},{rng('B')},$A{i},{rng('D')},$B{i})")
            c.number_format = NUM0
        pv.cell(row=i, column=11, value=f"=IF(E{i}>0,J{i}/E{i},\"\")").number_format = PCT
        for col, L in ((13, "X"), (14, "Y")):
            pv.cell(row=i, column=col, value=f"=SUMIFS({rng(L)},{rng('B')},$A{i},{rng('D')},$B{i})").number_format = NUM0
        pv.cell(row=i, column=15, value=f"=IF(M{i}>0,N{i}/M{i},\"\")").number_format = PCT
    pv.freeze_panes = "C5"

    nt = wb.create_sheet("How this works")
    for i, t in enumerate([
        "HOW THIS CLOSING WORKS",
        "Each sized set (from the sized-beam gate pass) is followed to the looms: beam length x mounting width multiplier (the ERP's own "
        "GET_BEAM_ALENGTH), and the fabric recorded against each beam (the ERP's own GET_BEAM_PROD).",
        "Length basis for weft: full beam length for beams taken off COMPLETE; fabric woven for beams CUT early, still on the loom, or not mounted.",
        "Weft needed = length basis x weft lb/m, where weft lb/m = picks per inch x width x 1.0936 / (840 x owner's weft count).",
        "Weft yarn available = yarn received from the party (net of returns) - yarn sent to sizing (net). Weft balance SHOULD BE = available - needed.",
        "Over / (under) issue = weft actually issued to looms - weft needed on the sets. Over-issue means more weft went to the looms "
        "than the sets could use (weft wasted, misplaced, or booked to this contract but used on another).",
        "Beam shortage % = 1 - fabric woven / set length (normal crimp and take-up is roughly 5-10%).",
        "Weft goes into each meter of FABRIC, and fabric is shorter than the set (crimp/take-up), so weft needed on the full set length is "
        "about the beam shortage % higher than weft needed on the fabric woven. Both are shown: set length (owner's basis) and fabric woven.",
        "Contracts started before Oct-2023: old-ERP yarn has no contract number, so their weft account is incomplete.",
        "Nothing is changed in the ERP.",
    ], 1):
        c = nt.cell(row=i, column=1, value=t)
        if t.isupper():
            c.font = Font(bold=True, color="1F3B57")
    nt.column_dimensions["A"].width = 150
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("contracts:", len(order), "sets:", len(sets), "tabs:", len(wb.sheetnames))


if __name__ == "__main__":
    main()
    if len(sys.argv) > 2 and sys.argv[1] == "--upload":
        cc.OUT = OUT
        cc.upload(sys.argv[2])


