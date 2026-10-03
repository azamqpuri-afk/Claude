"""
61_beam_sizer_audit.py - empty beams sent to each sizing party vs beams that came back (filled or empty), by COUNT
(beam numbers ignored), with the ERP booking errors corrected; plus beams whose number was sent and never seen again.
ERP read only.  python 61_beam_sizer_audit.py --upload <google sheet id>
"""
import importlib.util
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
_s = importlib.util.spec_from_file_location("d", HERE / "40_business_dashboard.py")
dash = importlib.util.module_from_spec(_s); _s.loader.exec_module(dash)
OUT = HERE / "financial_statements" / "beam_sizer_audit.xlsx"
NAVY = PatternFill("solid", fgColor="1F3B57"); W = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
T = Font(bold=True, size=15, color="1F3B57"); NOTE = Font(italic=True, color="4B4F5A")
RED = PatternFill("solid", fgColor="F8C9C9"); AMB = PatternFill("solid", fgColor="FFE3B3"); GRN = PatternFill("solid", fgColor="D9F2E3")
# corrections of ERP booking errors found in the audit
MOVE = {("OGP_EMPTY_BEAMS", 17, "2023-12-31"): "MANZOOR SIZING INDUSTRIES"}   # booked to Ali Hajwery; the same 4 beams came back from Manzoor next day


def head(ws, r, labels, widths):
    for i, (h, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=r, column=i, value=h); c.font, c.fill = W, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center"); ws.column_dimensions[c.column_letter].width = w


def main():
    cur = dash.connect(dash.load_config()).cursor()
    q = lambda s: cur.execute(s).fetchall()  # noqa: E731
    ev = q("""SELECT a.an, TRUNC(m.td), m.tt, m.tno, m.gp_no, bm.bcode, d.beam_id FROM dostex.wv_gp_beamsm m JOIN dostex.wv_gp_beamsd d ON d.trans_no=m.trans_no
              JOIN dostex.beams bm ON bm.beam_id=d.beam_id LEFT JOIN dostex.ach a ON a.ac=m.ac1
              UNION ALL SELECT a.an, TRUNC(m.td), m.tt, m.tno, m.gp_no, b.bcode, b.beam_id FROM dostex.wytm m JOIN dostex.wybeams b ON b.trans_no=m.trans_no
              LEFT JOIN dostex.ach a ON a.ac=NVL(m.ac1,m.ac2) WHERE m.tt='IGP_SIZED_BEAMS'""")
    erp_view = dict(q("SELECT an, SUM(beams) FROM dostex.wv_empty_beam_stock_v GROUP BY an"))
    raw = defaultdict(lambda: defaultdict(int)); rows = []; seen = set(); issues = []
    for an, td, tt, tno, gp, code, bid in ev:
        raw[an][tt] += 1
        key = (tt, tno, code)
        if tt == "OGP_EMPTY_BEAMS" and key in seen:
            issues.append((td.date(), an, f"Empty-beam gate pass {tno} (GP {gp}) lists beam {code} twice - counted once"))
            continue
        seen.add(key)
        mv = MOVE.get((tt, tno, str(td.date())))
        if mv:
            if code == min(c for a, t, tt2, n, g, c, b in ev if tt2 == tt and n == tno):
                issues.append((td.date(), an, f"Empty-beam gate pass {tno} (GP {gp}, 4 beams) booked to {an}; the same 4 beams came back from {mv} the next day - moved to {mv}"))
            an = mv
        rows.append((an, td, 0 if tt == "OGP_EMPTY_BEAMS" else 1, tt, tno, gp, code, bid))
    # beams whose number was sent and never seen in the mill again
    back = defaultdict(list)
    for bid, td in q("""SELECT b.beam_id, m.td FROM dostex.wybeams b JOIN dostex.wytm m ON m.trans_no=b.trans_no
                        UNION ALL SELECT d.beam_id, m.td FROM dostex.wv_gp_beamsd d JOIN dostex.wv_gp_beamsm m ON m.trans_no=d.trans_no WHERE m.tt='IGP_EMPTY_BEAMS'
                        UNION ALL SELECT d.beam_id, m.td FROM dostex.wprdtd d JOIN dostex.wprdtm m ON m.trans_no=d.trans_no WHERE m.tt IN ('BEAMS_TO_LOOM','BEAMS_FROM_LOOM')"""):
        back[bid].append(td)
    lost = defaultdict(list)
    for an, td, p, tt, tno, gp, code, bid in rows:
        if tt == "OGP_EMPTY_BEAMS" and not [x for x in back[bid] if x >= td]:
            lost[an].append((code, td.date(), gp))
    wb = Workbook(); sm = wb.active; sm.title = "Summary"
    sm["A1"] = f"EMPTY BEAMS AT SIZING PARTIES - count audit · {date.today():%d-%b-%Y}"; sm["A1"].font = T
    sm["A2"] = ("Sent = empty beams on OGP_EMPTY_BEAMS gate passes. Back = filled beams received (IGP_SIZED_BEAMS) + empty beams returned (IGP_EMPTY_BEAMS). "
                "Beam numbers ignored. ERP read only."); sm["A2"].font = NOTE
    head(sm, 4, ["Sizing party", "Sent", "Back filled", "Back empty", "Balance as entered", "Balance after correcting booking errors",
                 "Lowest running balance (date)", "Beams sent & never seen again (numbers)", "Last movement", "ERP's own empty-beam report says", "Should hold now if every entry is right",
                 "Verdict"], [34, 7, 8, 8, 10, 12, 16, 40, 11, 12, 10, 90])
    by = defaultdict(list)
    for r in rows:
        by[r[0]].append(r)
    ledger = []
    r = 5
    verdict_rows = []
    for an in sorted(by, key=str):
        bal = mn = 0; mnd = None
        for x in sorted(by[an], key=lambda z: (z[1], z[2])):
            bal += 1 if x[3] == "OGP_EMPTY_BEAMS" else -1
            if bal < mn:
                mn, mnd = bal, x[1]
            ledger.append((an, x[1].date(), x[3], x[4], x[5], x[6], bal))
        v = raw[an]
        entered = v["OGP_EMPTY_BEAMS"] - v["IGP_SIZED_BEAMS"] - v["IGP_EMPTY_BEAMS"]
        last = max(x[1] for x in by[an]).date()
        ls = lost.get(an, [])
        should = bal - min(mn, 0)           # beams the party must hold if every ERP entry is right (it can never hold fewer than 0)
        sends = [x for x in sorted(by[an], key=lambda z: (z[1], z[2])) if x[3] == "OGP_EMPTY_BEAMS"]
        recent = sends[-bal:] if bal > 0 else []
        if an.startswith("HIBA"):
            should = 0                        # Hiba's own beams - only received vs sent back matters
            verdict = (f"Hiba's own beams (owner): filled beams received {v['IGP_SIZED_BEAMS']}, empty beams sent back {v['OGP_EMPTY_BEAMS']} as entered "
                       f"({v['OGP_EMPTY_BEAMS'] - 1} if the gate pass listing 00002 twice is one beam) - SQUARED by count.")
        elif should == 0:
            verdict = "Squared - every beam sent came back."
        elif bal > 0 and recent and all((date.today() - x[1].date()).days <= 30 for x in recent) and should == bal - min(mn, 0) and should <= bal + 1:
            verdict = f"Holds {should} beams - sent {recent[-1][1]:%d-%b-%y}, being sized now (normal)."
        else:
            verdict = (f"CHECK: by count it should still hold {should} beams (entered balance {bal:+d}; it returned {-mn} beams with no send entry, lowest on "
                       f"{mnd:%d-%b-%y})." if mnd else f"CHECK: by count it should still hold {should} beams.")
            verdict += (" But every beam NUMBER sent to it came back to the mill" if not ls else
                        " Numbers sent and never seen again: " + ", ".join(f"{c} ({d:%d-%b-%y})" for c, d, g in ls)) + \
                       " - the difference is send/return entries missing in the ERP. Ask the party for its count of our beams."
        vals = [an, v["OGP_EMPTY_BEAMS"], v["IGP_SIZED_BEAMS"], v["IGP_EMPTY_BEAMS"], entered, bal, f"{mn} ({mnd:%d-%b-%y})" if mnd else "0",
                ", ".join(f"{c} ({d:%d-%b-%y})" for c, d, g in ls), last, erp_view.get(an), should, verdict]
        for i, val in enumerate(vals, 1):
            c = sm.cell(row=r, column=i, value=val); c.alignment = Alignment(wrap_text=True, vertical="top")
        sm.cell(row=r, column=9).number_format = "dd-mmm-yy"
        sm.cell(row=r, column=12).fill = RED if verdict.startswith("CHECK") else (AMB if "being sized" in verdict else GRN)
        verdict_rows.append((an, bal, should, verdict)); r += 1
    tot_s = sum(v["OGP_EMPTY_BEAMS"] for v in raw.values()); tot_b = sum(v["IGP_SIZED_BEAMS"] + v["IGP_EMPTY_BEAMS"] for v in raw.values())
    sm.cell(row=r, column=1, value="ALL PARTIES").font = B
    for i, val in [(2, tot_s), (5, tot_s - tot_b), (6, sum(x[1] for x in verdict_rows)), (11, sum(x[2] for x in verdict_rows))]:
        sm.cell(row=r, column=i, value=val).font = B
    notes = [
        "The ERP's own empty-beam report (WV_EMPTY_BEAM_STOCK_V) counts empty beams RETURNED by a sizer (IGP_EMPTY_BEAMS) as if they were sent (+1 instead of -1), "
        "so it overstates what Haji Taj Babu, Manzoor and FM Textile hold. Use the 'Balance after correcting' column.",
        "Beams 'sent & never seen again' by number mostly came back under another number (typing slips like 00223/00222, 00382/00383, 00330/00339, 00091/00081, 00061/00161) "
        "or were renumbered - the count, not the number, decides what a party holds.",
        "Hiba Fabrics delivered 16 filled beams in Jul-2025 before any empty beam was sent to it - these came on Hiba's own beams (numbered by the mill: 00002-00237); "
        "the mill then sent 17 empty beams to Hiba (one gate pass lists 00002 twice).",
    ]
    for i, n_ in enumerate(notes, r + 2):
        sm.cell(row=i, column=1, value="• " + n_).font = NOTE
    iss = wb.create_sheet("Booking errors")
    head(iss, 1, ["Date", "Party", "Issue"], [12, 34, 140])
    for x in sorted(issues):
        iss.append(list(x)); iss.cell(row=iss.max_row, column=1).number_format = "dd-mmm-yy"
    iss.append([None, "ERP report", "WV_EMPTY_BEAM_STOCK_V counts IGP_EMPTY_BEAMS (empty beams returned) as +1 instead of -1"])
    lg = wb.create_sheet("Ledger")
    head(lg, 1, ["Party", "Date", "Movement", "Entry no.", "Gate pass", "Beam", "Running balance"], [34, 12, 18, 9, 9, 9, 10])
    for x in ledger:
        lg.append(list(x)); lg.cell(row=lg.max_row, column=2).number_format = "dd-mmm-yy"
    lg.freeze_panes = "A2"; lg.auto_filter.ref = f"A1:G{lg.max_row}"
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    for x in verdict_rows:
        print(x)
    print("total sent", tot_s, "back", tot_b, "issues", len(issues))


def upload(fid):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    creds = service_account.Credentials.from_service_account_file(str(HERE / "service_account.json"), scopes=["https://www.googleapis.com/auth/drive"])
    build("drive", "v3", credentials=creds).files().update(fileId=fid, media_body=MediaFileUpload(
        str(OUT), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")).execute()
    print("uploaded", fid)


if __name__ == "__main__":
    main()
    if len(sys.argv) > 2 and sys.argv[1] == "--upload":
        upload(sys.argv[2])
